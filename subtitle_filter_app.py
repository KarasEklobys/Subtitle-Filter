"""
Subtitle Filter
---------------
Оставляет в субтитрах (.srt / .vtt / .ass / .ssa) только строки с нужными
словами (или наоборот — убирает их, в зависимости от режима).

Запуск с интерфейсом:  python subtitle_filter_app.py
Запуск без интерфейса: python subtitle_filter_app.py --cli --help
Сборка в .exe: см. README.md
"""

import os
import re
import csv
import sys
import glob
import json
import shutil
import argparse
import datetime
import tempfile
import subprocess
import webbrowser
import threading
import urllib.request
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP_TITLE = "Subtitle Filter"
APP_VERSION = "1.3.0"

# Замени на свои владельца/репозиторий, если хочешь, чтобы кнопка
# "Проверить обновления" смотрела на твой GitHub. Если оставить как есть
# и репозитория с таким именем нет — кнопка просто покажет ошибку сети.
GITHUB_OWNER = "KarasEklobys"
GITHUB_REPO = "Subtitle-Filter"
RELEASE_ASSET_NAME = "SubtitleFilter.exe"

ALL_GROUPS_LABEL = "🔗 Все группы"
DEFAULT_GROUP = "Основной"
NO_GROUP_OVERRIDE = "(использовать текущую)"

MODE_KEEP = "keep"
MODE_REMOVE = "remove"

# ---------- i18n ----------
# Interface language: "ru" (default) or "en". Loaded from settings at
# startup; changing it takes effect after restarting the app.
LANG = "ru"

TRANSLATIONS = {
    "🔗 Все группы": "🔗 All groups",
    "Основной": "Main",
    "(использовать текущую)": "(use current)",
    "папа": "example",
    "беги": "run",
    "зажигательн": "spoil",
    "кредит": "credit",
    "Пакетная обработка субтитров без графического интерфейса.": "Batch-process subtitles without the GUI.",
    "Файл со словами (по одному на строку)": "Word list file (one word per line)",
    "Файл со словами-исключениями (необязательно)": "Exclusion word list file (optional)",
    "Файлы и/или папки с .srt/.vtt/.ass/.ssa": "Files and/or folders with .srt/.vtt/.ass/.ssa",
    "keep — оставить совпадения (по умолчанию), remove — убрать их": "keep — keep matching lines (default), remove — remove them",
    "Искать только целые слова": "Match whole words only",
    "Учитывать регистр букв": "Case-sensitive",
    "Слова — регулярные выражения": "Treat words as regular expressions",
    "Нижняя граница времени HH:MM:SS": "Start time HH:MM:SS",
    "Верхняя граница времени HH:MM:SS": "End time HH:MM:SS",
    "Папка для результатов (по умолчанию — filtered рядом с первым файлом)": "Output folder (default: filtered next to the first file)",
    "Не найдено ни одного .srt/.vtt/.ass/.ssa файла по указанным путям.": "No .srt/.vtt/.ass/.ssa files found at the given paths.",
    "Список слов пуст.": "The word list is empty.",
    "Выбрать файлы...": "Select files...",
    "Выбрать папку...": "Select folder...",
    "Экспорт отчёта (CSV)...": "Export report (CSV)...",
    "Экспорт отчёта (Excel)...": "Export report (Excel)...",
    "Отчёт по всем группам...": "Report for all groups...",
    "Скопировать сводку в буфер": "Copy summary to clipboard",
    "Выход": "Exit",
    "Файл": "File",
    "Новая группа...": "New group...",
    "Переименовать группу...": "Rename group...",
    "Удалить группу": "Delete group",
    "Сохранить группу...": "Save group...",
    "Загрузить в группу...": "Load into group...",
    "Профили": "Profiles",
    "Правка": "Edit",
    "Светлая": "Light",
    "Тёмная": "Dark",
    "Как в системе": "System",
    "Переключить тему": "Toggle theme",
    "Вид": "View",
    "Анализ без сохранения": "Dry-run analysis",
    "История обработок...": "Processing history...",
    "Журнал": "Log",
    "Проверить обновления": "Check for updates",
    "О программе": "About",
    "Справка": "Help",
    "Сохранить текущие настройки как профиль...": "Save current settings as profile...",
    "Загрузить": "Load",
    "Удалить": "Delete",
    "Оставляет в субтитрах только то, что важно.": "Keeps only the subtitle lines that matter.",
    "Открыть репозиторий на GitHub": "Open repository on GitHub",
    "Закрыть": "Close",
    "перетаскивание файлов включено": "drag-and-drop enabled",
    "перетаскивание выключено (нет tkinterdnd2)": "drag-and-drop disabled (tkinterdnd2 missing)",
    "Слова": "Words",
    "Настройки поиска": "Search settings",
    "ГРУППЫ СЛОВ": "WORD GROUPS",
    "+ группа": "+ group",
    "Переименовать": "Rename",
    "СЛОВА В ГРУППЕ (по одному в строке)": "WORDS IN GROUP (one per line)",
    "0 слов": "0 words",
    "Загрузить...": "Load...",
    "Сохранить...": "Save...",
    "ИСКЛЮЧИТЬ СЛОВА (всегда вырезаются из результата)": "EXCLUDE WORDS (always cut from the result)",
    "НАСТРОЙКИ ПОИСКА": "SEARCH SETTINGS",
    "Только целые слова": "Whole words only",
    "Регулярные выражения (regex)": "Regular expressions (regex)",
    "Режим:": "Mode:",
    "Оставить": "Keep",
    "Убрать": "Remove",
    "ФИЛЬТР ПО ВРЕМЕНИ": "TIME FILTER",
    "Ограничить диапазон времени": "Limit to a time range",
    "От": "From",
    "До": "To",
    "Формат: ЧЧ:ММ:СС, например 00:05:00": "Format: HH:MM:SS, e.g. 00:05:00",
    "ФАЙЛЫ СУБТИТРОВ (.srt / .vtt / .ass / .ssa)": "SUBTITLE FILES (.srt / .vtt / .ass / .ssa)",
    "Выбрать папку": "Select folder",
    "Удалить выбранное": "Remove selected",
    "Очистить всё": "Clear all",
    "Файлов не выбрано (можно перетащить сюда)": "No files selected (you can drag & drop here)",
    "Обработать": "Process",
    "Открыть папку с результатом": "Open result folder",
    "Скопировать сводку": "Copy summary",
    "РЕЗУЛЬТАТ ОБРАБОТКИ": "PROCESSING RESULT",
    "ПРЕДПРОСМОТР (двойной клик по строке — открыть видео рядом на этом месте)": "PREVIEW (double-click a line to open the matching video at that spot)",
    "Новая группа": "New group",
    "Название новой группы слов:": "New word group name:",
    "Группа с таким названием уже есть.": "A group with that name already exists.",
    "«Все группы» — это не отдельная группа, переименовать нельзя.": "“All groups” isn't a real group — it can't be renamed.",
    "Переименовать группу": "Rename group",
    "Новое название:": "New name:",
    "Нельзя удалить единственную группу.": "Can't delete the only group.",
    "ОК": "OK",
    "Отмена": "Cancel",
    "«Убрать» — из субтитров удаляются строки с совпадениями, всё остальное остаётся.": "“Remove” — lines that match get cut, everything else stays.",
    "«Оставить» — в субтитрах остаются только строки с совпадениями (обычный режим).": "“Keep” — only matching lines stay (the default mode).",
    "Учитывает персональную группу файла (если назначена через контекстное меню).": "Takes into account a file's own group, if one was assigned via the right-click menu.",
    "Выбери конкретную группу, чтобы загрузить в неё список.": "Pick a specific group to load the list into.",
    "Загрузить список слов": "Load word list",
    "Текстовый файл": "Text file",
    "Сохранить список слов": "Save word list",
    "Новый профиль": "New profile",
    "Название профиля:": "Profile name:",
    "Выбери файлы субтитров": "Select subtitle files",
    "Субтитры": "Subtitles",
    "Выбери папку с субтитрами": "Select a folder with subtitles",
    "Назначить группу для выбранных...": "Assign group to selected...",
    "Сбросить группу для выбранных": "Reset group for selected",
    "Удалить из списка": "Remove from list",
    "Группа для файлов": "Group for files",
    "Использовать группу:": "Use group:",
    "Выбери файл в списке слева.": "Select a file from the list on the left.",
    "Совпадений не найдено.": "No matches found.",
    "Видео с таким же именем рядом с субтитрами не найдено.": "No video with a matching name found next to the subtitles.",
    "VLC не найден — видео открыто с начала (перемотка недоступна).": "VLC wasn't found — the video opened from the start (can't seek).",
    "Сначала выбери файлы субтитров.": "Select subtitle files first.",
    "Анализ...\n": "Analyzing...\n",
    "Чаще всего встречалось: ": "Most frequent: ",
    "Анализ завершён — файлы не изменены.": "Analysis complete — no files were changed.",
    "Сначала добавь хотя бы одно слово в список.": "Add at least one word to the list first.",
    "Обработка...\n": "Processing...\n",
    "для этого файла список слов пуст": "the word list is empty for this file",
    "История обработок": "Processing history",
    "Пока пусто — здесь появится история после первой обработки.": "Nothing here yet — history will show up after the first run.",
    "Папка результата больше не существует.": "The result folder doesn't exist anymore.",
    "Двойной клик — открыть папку с результатом этого запуска.": "Double-click to open the result folder for that run.",
    "Сначала обработай файлы.": "Process files first.",
    "Сохранить отчёт": "Save report",
    "Слово": "Word",
    "Найдено раз": "Matches",
    "Для экспорта в Excel нужна библиотека openpyxl:\npip install openpyxl": "Exporting to Excel needs the openpyxl library:\npip install openpyxl",
    "Отчёт": "Report",
    "Формат отчёта": "Report format",
    "В каком формате сохранить:": "Save as:",
    "Совпадений не найдено ни в одной группе.": "No matches found in any group.",
    "Сохранить отчёт по всем группам": "Save report for all groups",
    "Все группы": "All groups",
    "Группа": "Group",
    "По словам": "By word",
    "Всего": "Total",
    "Найдено раз по словам": "Matches by word",
    "Раз": "Matches",
    "Сводка скопирована в буфер обмена.": "Summary copied to clipboard.",
    "У тебя установлена последняя версия.": "You're already on the latest version.",
    "Скачиваю обновление...": "Downloading update...",
    "Неверное регулярное выражение: ": "Invalid regular expression: ",
    " из ": " of ",
    " строк": " lines",
    "ОШИБКА ": "ERROR ",
    "\nИтого: ": "\nTotal: ",
    " строк в ": " lines in ",
    " файл(ах). Результаты в: ": " file(s). Results in: ",
    "Версия ": "Version ",
    "Оставляет в .srt/.vtt/.ass/.ssa только строки с нужными словами (или убирает их). ": "Keeps only the lines you need in .srt/.vtt/.ass/.ssa files (or removes them). ",
    "» объединяет слова из всех групп сразу.": "” combines words from every group at once.",
    "Удалить группу «": "Delete group “",
    " слов(а)": " word(s)",
    "Профиль «": "Profile “",
    "» сохранён.": "” saved.",
    "» загружен.": "” loaded.",
    "Удалить профиль «": "Delete profile “",
    "Ничего не найдено по «": "Nothing found for “",
    " файл(ов) всего)": " file(s) total)",
    " файл(ов)": " file(s)",
    " файл(ов) выбрано": " file(s) selected",
    "Не удалось прочитать файл: ": "Couldn't read the file: ",
    "... и ещё ": "... and ",
    " строк (показаны первые 200)": " more lines (first 200 shown)",
    "Не удалось открыть видео: ": "Couldn't open the video: ",
    " — ошибка: ": " — error: ",
    " — совпало бы ": " — would match ",
    "\nИтого без сохранения: ": "\nTotal (dry run): ",
    " строк.": " lines.",
    " — 0 из ": " — 0 of ",
    " строк (совпадений нет)": " lines (no matches)",
    " файл(ах).": " file(s).",
    "\nГотово. Результаты в папке:\n": "\nDone. Results in:\n",
    "Готово: ": "Done: ",
    " файл(ов)  ·  ": " file(s)  ·  ",
    "Отчёт сохранён: ": "Report saved: ",
    "Отчёт по всем группам сохранён: ": "Report for all groups saved: ",
    "Не удалось проверить обновления.\n(": "Couldn't check for updates.\n(",
    "Доступна новая версия: ": "A new version is available: ",
    " (у тебя ": " (you have ",
    ").\n\nДа — скачать и установить автоматически.\nНет — открыть страницу релиза в браузере.": ").\n\nYes — download and install automatically.\nNo — open the release page in your browser.",
    ").\nОткрыть страницу релиза?": ").\nOpen the release page?",
    "Не удалось установить обновление:\n": "Couldn't install the update:\n"
}

def t(s):
    """Translate a UI string to the current LANG (English only; Russian is the source language)."""
    if LANG == "en":
        return TRANSLATIONS.get(s, s)
    return s

SUPPORTED_EXTS = (".srt", ".vtt", ".ass", ".ssa")

# ---------- темы (в духе macOS: светлая/тёмная/системная) ----------

THEMES = {
    "light": dict(
        bg="#f0f0f3", panel="#ffffff", panel2="#f5f5f7", border="#dcdce1",
        text="#1d1d1f", text_dim="#6e6e73", accent="#0071e3", accent_active="#0060c2",
        accent_text="#ffffff", danger="#d70015", hit="#0071e3", track_off="#d5d5da",
    ),
    "dark": dict(
        bg="#1e1e1e", panel="#262626", panel2="#2c2c2e", border="#3a3a3c",
        text="#f5f5f7", text_dim="#98989d", accent="#0a84ff", accent_active="#3aa0ff",
        accent_text="#ffffff", danger="#ff453a", hit="#5ec2a0", track_off="#48484a",
    ),
}

FONT_UI = ("Segoe UI", 10)
FONT_UI_DIM = ("Segoe UI", 9)
FONT_UI_BOLD_SM = ("Segoe UI", 9, "bold")
FONT_MONO = ("Consolas", 9)
FONT_MONO_KW = ("Consolas", 10)

# опциональный drag-and-drop: работает, только если установлена tkinterdnd2
try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    BaseTk = TkinterDnD.Tk
    DND_AVAILABLE = True
except ImportError:
    BaseTk = tk.Tk
    DND_AVAILABLE = False

# опциональный экспорт в Excel — работает, только если установлен openpyxl
try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def detect_system_theme():
    if os.name != "nt":
        return "light"
    try:
        import winreg
        key = winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize",
        )
        value, _ = winreg.QueryValueEx(key, "AppsUseLightTheme")
        return "light" if value else "dark"
    except Exception:
        return "light"


# ---------- пути / настройки ----------

def app_dir():
    """Папка рядом с .exe (или рядом со скриптом, если запущено как .py)."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def settings_dir():
    if os.name == "nt":
        base = os.environ.get("APPDATA") or app_dir()
        d = os.path.join(base, "SubtitleFilter")
    else:
        d = os.path.join(os.path.expanduser("~"), ".subtitle_filter")
    try:
        os.makedirs(d, exist_ok=True)
    except Exception:
        return app_dir()
    return d


SETTINGS_PATH = os.path.join(settings_dir(), "settings.json")
_OLD_SETTINGS_PATH = os.path.join(app_dir(), "settings.json")

DEFAULT_SETTINGS = {
    "whole_word": True,
    "case_sensitive": False,
    "regex_mode": False,
    "mode": MODE_KEEP,
    "theme": "system",
    "lang": "ru",
    "last_folder": "",
    "active_group": DEFAULT_GROUP,
    "groups": {DEFAULT_GROUP: ["папа", "беги", "зажигательн", "кредит"]},
    "exclude_words": [],
    "time_filter_enabled": False,
    "time_filter_start": "",
    "time_filter_end": "",
    "file_groups": {},
    "profiles": {},
    "history": [],
    "window_geometry": "1150x760",
}


def load_settings():
    if not os.path.exists(SETTINGS_PATH) and os.path.exists(_OLD_SETTINGS_PATH):
        try:
            shutil.copy(_OLD_SETTINGS_PATH, SETTINGS_PATH)
        except Exception:
            pass
    if os.path.exists(SETTINGS_PATH):
        try:
            with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            merged = dict(DEFAULT_SETTINGS)
            merged.update(data)
            if not merged.get("groups"):
                merged["groups"] = dict(DEFAULT_SETTINGS["groups"])
            return merged
        except Exception:
            pass
    return json.loads(json.dumps(DEFAULT_SETTINGS))


def save_settings(settings):
    try:
        with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
    except Exception:
        pass  # настройки — не критично, если не удалось сохранить


# ---------- время ----------

def time_to_seconds(t):
    """Понимает и '00:01:02,003' (srt/vtt), и '0:01:02.03' (ass/ssa), и просто 'HH:MM:SS'."""
    if not t:
        return None
    m = re.search(r"(\d+):(\d{2}):(\d{2})(?:[.,](\d+))?", t)
    if not m:
        return None
    h, mi, s, frac = m.groups()
    frac_ms = ((frac or "0") + "000")[:3]
    try:
        return int(h) * 3600 + int(mi) * 60 + int(s) + int(frac_ms) / 1000.0
    except ValueError:
        return None


# ---------- парсинг субтитров (.srt, .vtt, .ass/.ssa) ----------

def parse_srt(text: str):
    text = text.replace("\r\n", "\n").lstrip("\ufeff")
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    parsed = []
    for block in blocks:
        lines = block.split("\n")
        if len(lines) < 2:
            continue
        time_idx = 0 if "-->" in lines[0] else 1
        if time_idx >= len(lines) or "-->" not in lines[time_idx]:
            continue
        time_line = lines[time_idx]
        body = "\n".join(lines[time_idx + 1:])
        parsed.append({"time": time_line, "text": body})
    return parsed


def rebuild_srt(blocks):
    out = []
    for i, b in enumerate(blocks, start=1):
        out.append(f"{i}\n{b['time']}\n{b['text']}")
    return ("\n\n".join(out) + "\n") if out else ""


def parse_vtt(text: str):
    text = text.replace("\r\n", "\n").lstrip("\ufeff")
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    parsed = []
    for block in blocks:
        if block.upper().startswith("WEBVTT"):
            continue
        lines = block.split("\n")
        time_idx = None
        for i, line in enumerate(lines):
            if "-->" in line:
                time_idx = i
                break
        if time_idx is None:
            continue
        time_line = lines[time_idx]
        body = "\n".join(lines[time_idx + 1:])
        if body:
            parsed.append({"time": time_line, "text": body})
    return parsed


def rebuild_vtt(blocks):
    out = ["WEBVTT"]
    for b in blocks:
        out.append(f"{b['time']}\n{b['text']}")
    return "\n\n".join(out) + "\n"


def parse_subtitle(text: str, ext: str):
    if ext == ".vtt":
        return parse_vtt(text)
    return parse_srt(text)


def rebuild_subtitle(blocks, ext: str):
    if ext == ".vtt":
        return rebuild_vtt(blocks)
    return rebuild_srt(blocks)


def strip_tags(s: str) -> str:
    return re.sub(r"<[^>]+>", "", s)


def strip_ass_tags(s: str) -> str:
    s = re.sub(r"\{[^}]*\}", "", s)
    return s.replace("\\N", "\n").replace("\\n", "\n").replace("\\h", " ")


def parse_ass_events(lines):
    """Находит секцию [Events] и разбирает строки Dialogue: по формату из Format:."""
    events_start = None
    for i, line in enumerate(lines):
        if line.strip().lower() == "[events]":
            events_start = i
            break
    if events_start is None:
        return []

    fmt_fields = None
    events = []
    for i in range(events_start + 1, len(lines)):
        line = lines[i]
        s = line.strip()
        if s.lower().startswith("format:"):
            fmt_fields = [f.strip().lower() for f in s[len("format:"):].split(",")]
            continue
        if s.lower().startswith("[") and s.endswith("]"):
            break  # новая секция — события закончились
        if s.lower().startswith("dialogue:") and fmt_fields:
            rest = line.split(":", 1)[1]
            values = rest.split(",", len(fmt_fields) - 1)
            if len(values) < len(fmt_fields):
                continue
            row = dict(zip(fmt_fields, values))
            events.append({
                "line_index": i,
                "start": row.get("start", "").strip(),
                "end": row.get("end", "").strip(),
                "text": row.get("text", ""),
            })
    return events


def filter_ass_lines(lines, decide_fn):
    events = parse_ass_events(lines)
    dialogue_indices = {e["line_index"] for e in events}
    kept_indices = set()
    for e in events:
        clean = strip_ass_tags(e["text"])
        start_sec = time_to_seconds(e["start"])
        if decide_fn(clean, start_sec):
            kept_indices.add(e["line_index"])
    out_lines = [ln for i, ln in enumerate(lines) if i not in dialogue_indices or i in kept_indices]
    kept_count = len(kept_indices)
    total_count = len(events)
    plain_text = "\n".join(strip_ass_tags(e["text"]) for e in events)
    return out_lines, kept_count, total_count, plain_text


# ---------- поиск по словам ----------

def escape_and_sort(keywords):
    return sorted((re.escape(k) for k in keywords), key=len, reverse=True)


def build_matcher(keywords, whole_word: bool, case_sensitive: bool = False, regex_mode: bool = False):
    """Может бросить re.error, если regex_mode=True и шаблон некорректен."""
    if not keywords:
        return None
    flags = re.UNICODE | (0 if case_sensitive else re.IGNORECASE)
    if regex_mode:
        alt = "|".join(f"(?:{k})" for k in keywords)
        return re.compile(f"(?:{alt})", flags)

    alt = "|".join(escape_and_sort(keywords))
    if whole_word:
        pattern = rf"(?<![^\W\d_])(?:{alt})(?![^\W\d_])"
        try:
            return re.compile(pattern, flags)
        except re.error:
            pass
    return re.compile(f"(?:{alt})", flags)


def build_single_matcher(keyword, whole_word: bool, case_sensitive: bool = False, regex_mode: bool = False):
    flags = re.UNICODE | (0 if case_sensitive else re.IGNORECASE)
    core = keyword if regex_mode else re.escape(keyword)
    if whole_word and not regex_mode:
        pattern = rf"(?<![^\W\d_]){core}(?![^\W\d_])"
        try:
            return re.compile(pattern, flags)
        except re.error:
            pass
    return re.compile(core, flags)


def decide_keep(text, start_sec, matcher, invert, exclude_matcher,
                 time_enabled=False, time_start=None, time_end=None):
    """Единая логика решения «оставить строку или нет» — используется и в GUI, и в CLI."""
    if time_enabled:
        if start_sec is not None:
            if time_start is not None and start_sec < time_start:
                return False
            if time_end is not None and start_sec > time_end:
                return False
    matched = bool(matcher.search(text)) if matcher else False
    keep = (not matched) if invert else matched
    if not keep:
        return False
    if exclude_matcher and exclude_matcher.search(text):
        return False
    return True


def peek_subtitle_blocks(path):
    """Единое представление содержимого файла для предпросмотра/анализа:
    (ext, [{"start_sec": float|None, "text": str}, ...])."""
    ext = os.path.splitext(path)[1].lower()
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        content = f.read()
    if ext in (".ass", ".ssa"):
        lines = content.replace("\r\n", "\n").split("\n")
        events = parse_ass_events(lines)
        return ext, [
            {"start_sec": time_to_seconds(e["start"]), "text": strip_ass_tags(e["text"])}
            for e in events
        ]
    blocks = parse_subtitle(content, ext)
    out = []
    for b in blocks:
        start = b["time"].split("-->")[0].strip()
        out.append({"start_sec": time_to_seconds(start), "text": strip_tags(b["text"])})
    return ext, out


def filter_subtitle_file(path: str, decide_fn, out_dir: str, suffix: str = "filtered"):
    """decide_fn(text, start_sec) -> bool. Возвращает (out_path, kept_count, total_count, all_plain_text)."""
    ext = os.path.splitext(path)[1].lower()
    base = os.path.basename(path)
    name, _ = os.path.splitext(base)
    out_name = f"{name}.{suffix}{ext}"
    out_path = os.path.join(out_dir, out_name)

    if ext in (".ass", ".ssa"):
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            content = f.read()
        lines = content.replace("\r\n", "\n").split("\n")
        out_lines, kept_count, total_count, plain_text = filter_ass_lines(lines, decide_fn)
        out_text = "\n".join(out_lines)
        if not out_text.endswith("\n"):
            out_text += "\n"
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            content = f.read()
        blocks = parse_subtitle(content, ext)
        kept = []
        plain_parts = []
        for b in blocks:
            clean = strip_tags(b["text"])
            plain_parts.append(clean)
            start_sec = time_to_seconds(b["time"].split("-->")[0].strip())
            if decide_fn(clean, start_sec):
                kept.append(b)
        out_text = rebuild_subtitle(kept, ext)
        kept_count, total_count = len(kept), len(blocks)
        plain_text = "\n".join(plain_parts)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(out_text)

    return out_path, kept_count, total_count, plain_text


# ---------- CLI (пакетная обработка без интерфейса) ----------

def _gather_cli_files(inputs):
    files = []
    for item in inputs:
        if os.path.isdir(item):
            for ext in SUPPORTED_EXTS:
                files.extend(sorted(glob.glob(os.path.join(item, f"*{ext}"))))
        elif os.path.isfile(item):
            files.append(item)
    seen, unique = set(), []
    for f in files:
        if f not in seen:
            seen.add(f)
            unique.append(f)
    return unique


def run_cli(argv):
    global LANG
    LANG = load_settings().get("lang", "ru")
    parser = argparse.ArgumentParser(
        prog="SubtitleFilter --cli",
        description=t("Пакетная обработка субтитров без графического интерфейса."),
    )
    parser.add_argument("--cli", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--keywords", required=True, help=t("Файл со словами (по одному на строку)"))
    parser.add_argument("--exclude", help=t("Файл со словами-исключениями (необязательно)"))
    parser.add_argument("--input", nargs="+", required=True, help=t("Файлы и/или папки с .srt/.vtt/.ass/.ssa"))
    parser.add_argument("--mode", choices=[MODE_KEEP, MODE_REMOVE], default=MODE_KEEP,
                         help=t("keep — оставить совпадения (по умолчанию), remove — убрать их"))
    parser.add_argument("--whole-word", action="store_true", help=t("Искать только целые слова"))
    parser.add_argument("--case-sensitive", action="store_true", help=t("Учитывать регистр букв"))
    parser.add_argument("--regex", action="store_true", help=t("Слова — регулярные выражения"))
    parser.add_argument("--time-start", help=t("Нижняя граница времени HH:MM:SS"))
    parser.add_argument("--time-end", help=t("Верхняя граница времени HH:MM:SS"))
    parser.add_argument("--out", help=t("Папка для результатов (по умолчанию — filtered рядом с первым файлом)"))
    args = parser.parse_args(argv)

    with open(args.keywords, "r", encoding="utf-8-sig", errors="replace") as f:
        keywords = [w.strip() for w in f.read().splitlines() if w.strip()]
    exclude_words = []
    if args.exclude:
        with open(args.exclude, "r", encoding="utf-8-sig", errors="replace") as f:
            exclude_words = [w.strip() for w in f.read().splitlines() if w.strip()]

    files = _gather_cli_files(args.input)
    if not files:
        print(t("Не найдено ни одного .srt/.vtt/.ass/.ssa файла по указанным путям."))
        return 1
    if not keywords:
        print(t("Список слов пуст."))
        return 1

    try:
        matcher = build_matcher(keywords, args.whole_word, args.case_sensitive, args.regex)
        exclude_matcher = build_matcher(exclude_words, args.whole_word, args.case_sensitive, args.regex)
    except re.error as e:
        print((t('Неверное регулярное выражение: ') + f"{e}"))
        return 1

    time_enabled = bool(args.time_start or args.time_end)
    t_start = time_to_seconds(args.time_start) if args.time_start else None
    t_end = time_to_seconds(args.time_end) if args.time_end else None
    invert = args.mode == MODE_REMOVE
    suffix = "cleaned" if invert else "filtered"

    def decide(text, start_sec):
        return decide_keep(text, start_sec, matcher, invert, exclude_matcher, time_enabled, t_start, t_end)

    out_dir = args.out or os.path.join(os.path.dirname(os.path.abspath(files[0])), "filtered")
    os.makedirs(out_dir, exist_ok=True)

    exit_code = 0
    total_kept = total_all = 0
    for path in files:
        try:
            out_path, kept, total, _ = filter_subtitle_file(path, decide, out_dir, suffix)
            print((f"{os.path.basename(out_path)}" + ' — ' + f"{kept}" + t(' из ') + f"{total}" + t(' строк')))
            total_kept += kept
            total_all += total
        except Exception as e:
            print((t('ОШИБКА ') + f"{os.path.basename(path)}" + ': ' + f"{e}"))
            exit_code = 1

    print((t('\nИтого: ') + f"{total_kept}" + t(' из ') + f"{total_all}" + t(' строк в ') + f"{len(files)}" + t(' файл(ах). Результаты в: ') + f"{out_dir}"))
    return exit_code


# ---------- кастомные виджеты в духе macOS ----------

def _round_rect_points(x1, y1, x2, y2, r):
    r = max(0, min(r, (x2 - x1) / 2, (y2 - y1) / 2))
    return [
        x1 + r, y1, x2 - r, y1, x2, y1, x2, y1 + r,
        x2, y2 - r, x2, y2, x2 - r, y2, x1 + r, y2,
        x1, y2, x1, y2 - r, x1, y1 + r, x1, y1,
    ]


class ToggleSwitch(tk.Canvas):
    """Переключатель в стиле macOS вместо обычного чекбокса."""

    def __init__(self, parent, app, variable, command=None, width=42, height=24):
        super().__init__(parent, width=width, height=height, highlightthickness=0, bd=0, cursor="hand2")
        self.app = app
        self.var = variable
        self.command = command
        self.w = width
        self.h = height
        self.enabled = True
        self.bind("<Button-1>", self._on_click)
        app.register_themed(self)
        self.redraw()

    def set_enabled(self, enabled):
        self.enabled = enabled
        self.configure(cursor="hand2" if enabled else "arrow")
        self.redraw()

    def _on_click(self, event=None):
        if not self.enabled:
            return
        self.var.set(not self.var.get())
        self.redraw()
        if self.command:
            self.command()

    def redraw(self):
        self.delete("all")
        theme = self.app.theme
        self.configure(bg=theme["panel"])
        on = bool(self.var.get())
        track = theme["accent"] if on else theme["track_off"]
        if not self.enabled:
            track = theme["border"]
        r = self.h / 2
        self.create_polygon(_round_rect_points(1, 1, self.w - 1, self.h - 1, r - 1),
                             smooth=True, fill=track, outline="")
        knob_d = self.h - 6
        x = (self.w - knob_d - 3) if on else 3
        self.create_oval(x, 3, x + knob_d, 3 + knob_d, fill="#ffffff", outline="")


class SegmentedControl(tk.Canvas):
    """Сегментированный переключатель (как в macOS) для выбора из нескольких вариантов."""

    def __init__(self, parent, app, options, variable, command=None, width=240, height=28):
        super().__init__(parent, width=width, height=height, highlightthickness=0, bd=0, cursor="hand2")
        self.app = app
        self.options = options  # [(value, label), ...]
        self.var = variable
        self.command = command
        self.w = width
        self.h = height
        self.bind("<Button-1>", self._on_click)
        app.register_themed(self)
        self.redraw()

    def _segment_at(self, x):
        seg_w = self.w / len(self.options)
        idx = int(x // seg_w)
        return max(0, min(idx, len(self.options) - 1))

    def _on_click(self, event):
        idx = self._segment_at(event.x)
        value = self.options[idx][0]
        if value != self.var.get():
            self.var.set(value)
            self.redraw()
            if self.command:
                self.command()

    def redraw(self):
        self.delete("all")
        theme = self.app.theme
        self.configure(bg=theme["panel"])
        r = self.h / 2
        self.create_polygon(_round_rect_points(1, 1, self.w - 1, self.h - 1, r - 1),
                             smooth=True, fill=theme["panel2"], outline=theme["border"])
        seg_w = self.w / len(self.options)
        active_idx = next((i for i, (v, _) in enumerate(self.options) if v == self.var.get()), 0)
        pad = 3
        x1 = active_idx * seg_w + pad
        x2 = (active_idx + 1) * seg_w - pad
        self.create_polygon(_round_rect_points(x1, pad, x2, self.h - pad, (self.h - 2 * pad) / 2),
                             smooth=True, fill=theme["accent"], outline="")
        for i, (value, label) in enumerate(self.options):
            cx = i * seg_w + seg_w / 2
            fg = theme["accent_text"] if i == active_idx else theme["text_dim"]
            self.create_text(cx, self.h / 2, text=label, fill=fg, font=FONT_UI_DIM)


class PillButton(tk.Canvas):
    """Крупная скруглённая кнопка в духе macOS для главного действия."""

    def __init__(self, parent, app, text, command, width=170, height=38):
        super().__init__(parent, width=width, height=height, highlightthickness=0, bd=0, cursor="hand2")
        self.app = app
        self.text = text
        self.command = command
        self.w = width
        self.h = height
        self.enabled = True
        self.hover = False
        self.bind("<Button-1>", self._on_click)
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        app.register_themed(self)
        self.redraw()

    def set_enabled(self, enabled):
        self.enabled = enabled
        self.configure(cursor="hand2" if enabled else "arrow")
        self.redraw()

    def _on_enter(self, event=None):
        if self.enabled:
            self.hover = True
            self.redraw()

    def _on_leave(self, event=None):
        self.hover = False
        self.redraw()

    def _on_click(self, event=None):
        if self.enabled and self.command:
            self.command()

    def redraw(self):
        self.delete("all")
        theme = self.app.theme
        self.configure(bg=theme["panel"])
        if not self.enabled:
            fill = theme["track_off"]
            fg = theme["text_dim"]
        else:
            fill = theme["accent_active"] if self.hover else theme["accent"]
            fg = theme["accent_text"]
        self.create_polygon(_round_rect_points(0, 0, self.w, self.h, self.h / 2), smooth=True, fill=fill, outline="")
        self.create_text(self.w / 2, self.h / 2, text=self.text, fill=fg, font=("Segoe UI", 11, "bold"))


class RoundedPanel(tk.Frame):
    """Панель со скруглёнными углами (canvas-подложка + вложенный content-фрейм)."""

    def __init__(self, parent, app, radius=12):
        super().__init__(parent)
        self.app = app
        self.radius = radius
        self.canvas = tk.Canvas(self, highlightthickness=0, bd=0)
        self.canvas.pack(fill="both", expand=True)
        self.content = tk.Frame(self.canvas)
        self._win_id = self.canvas.create_window(radius, radius, window=self.content, anchor="nw")
        self.canvas.bind("<Configure>", self._on_resize)
        app.register_themed(self)

    def _on_resize(self, event=None):
        w = event.width if event else self.canvas.winfo_width()
        h = event.height if event else self.canvas.winfo_height()
        if w < 4 or h < 4:
            return
        r = self.radius
        self.canvas.delete("bg")
        pts = _round_rect_points(1, 1, w - 1, h - 1, r)
        self.canvas.create_polygon(pts, smooth=True, fill=self.app.theme["panel"],
                                    outline=self.app.theme["border"], tags="bg")
        self.canvas.tag_lower("bg")
        self.canvas.itemconfig(self._win_id, width=max(w - 2 * r, 10), height=max(h - 2 * r, 10))

    def redraw(self):
        self.canvas.configure(bg=self.app.theme["bg"])
        self.content.configure(bg=self.app.theme["panel"])
        self._on_resize()


class Toast(tk.Toplevel):
    """Всплывающее ненавязчивое уведомление вместо блокирующего messagebox."""

    def __init__(self, app, message, kind="info", duration=2800):
        super().__init__(app)
        self.overrideredirect(True)
        try:
            self.attributes("-topmost", True)
        except tk.TclError:
            pass
        t = app.theme
        bar_color = t["danger"] if kind in ("warn", "error") else t["accent"]
        outer = tk.Frame(self, bg=t["border"])
        outer.pack(fill="both", expand=True)
        inner = tk.Frame(outer, bg=t["panel"])
        inner.pack(fill="both", expand=True, padx=1, pady=1)
        tk.Frame(inner, bg=bar_color, width=4).pack(side="left", fill="y")
        tk.Label(inner, text=message, bg=t["panel"], fg=t["text"], font=FONT_UI,
                 wraplength=320, justify="left", padx=12, pady=10).pack(side="left")
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        try:
            x = app.winfo_x() + app.winfo_width() - w - 24
            y = app.winfo_y() + app.winfo_height() - h - 24
        except tk.TclError:
            x, y = 100, 100
        self.geometry(f"{w}x{h}+{max(x, 0)}+{max(y, 0)}")
        self.after(duration, self._close)

    def _close(self):
        try:
            self.destroy()
        except tk.TclError:
            pass


# ---------- интерфейс ----------

class App(BaseTk):
    def __init__(self):
        super().__init__()
        self.settings = load_settings()
        global LANG
        LANG = self.settings.get("lang", "ru")
        self.theme = THEMES.get(self._effective_theme_name(), THEMES["light"])

        self.title(f"{APP_TITLE} v{APP_VERSION}")
        self.geometry(self.settings.get("window_geometry", "1150x760"))
        self.minsize(940, 640)
        self._set_window_icon()

        self.files = []             # list of full paths
        self._visible_indices = []  # индексы self.files, отображаемые в listbox сейчас
        self.last_report = []       # [{file, keyword, count}, ...] после обработки
        self.last_out_dir = None
        self._preview_job = None
        self._themed_widgets = []   # canvas-виджеты, умеющие сами перекрашиваться (.redraw)
        self._raw_frames = []       # (widget, role) для tk.Frame с ручной темой
        self._raw_texts = []        # tk.Text / tk.Listbox / tk.Entry с ручной темой

        self.whole_word_var = tk.BooleanVar(value=self.settings.get("whole_word", True))
        self.case_sensitive_var = tk.BooleanVar(value=self.settings.get("case_sensitive", False))
        self.regex_mode_var = tk.BooleanVar(value=self.settings.get("regex_mode", False))
        self.mode_var = tk.StringVar(value=self.settings.get("mode", MODE_KEEP))
        self.theme_var = tk.StringVar(value=self.settings.get("theme", "system"))
        self.time_filter_var = tk.BooleanVar(value=self.settings.get("time_filter_enabled", False))
        self.time_start_var = tk.StringVar(value=self.settings.get("time_filter_start", ""))
        self.time_end_var = tk.StringVar(value=self.settings.get("time_filter_end", ""))
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *a: self._refresh_file_list())

        self._build_style()
        self._build_menu()
        self._build_ui()
        self._apply_theme()
        self._load_group_names()
        self._select_group(self.settings.get("active_group", DEFAULT_GROUP))
        self._load_exclude_words()
        self._sync_control_states()

        last_folder = self.settings.get("last_folder", "")
        if last_folder and os.path.isdir(last_folder):
            found = []
            for ext in SUPPORTED_EXTS:
                found.extend(glob.glob(os.path.join(last_folder, f"*{ext}")))
            self.files = sorted(found)
            self._refresh_file_list()

        if DND_AVAILABLE:
            self.file_listbox.drop_target_register(DND_FILES)
            self.file_listbox.dnd_bind("<<Drop>>", self._on_drop_files)

        self._bind_shortcuts()
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _effective_theme_name(self):
        name = self.settings.get("theme", "system")
        if name == "system":
            return detect_system_theme()
        return name if name in THEMES else "light"

    def register_themed(self, widget):
        self._themed_widgets.append(widget)

    def _set_window_icon(self):
        icon_path = os.path.join(app_dir(), "app_icon.ico")
        if os.path.exists(icon_path):
            try:
                self.iconbitmap(icon_path)
            except Exception:
                pass

    def show_toast(self, message, kind="info"):
        try:
            Toast(self, message, kind=kind)
        except tk.TclError:
            pass

    # ---------- стиль ----------

    def _build_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        self.ttk_style = style

    def _apply_theme(self):
        t = self.theme
        style = self.ttk_style

        self.configure(bg=t["bg"])
        style.configure("TFrame", background=t["bg"])
        style.configure("Panel.TFrame", background=t["panel"])
        style.configure("TLabel", background=t["bg"], foreground=t["text"], font=FONT_UI)
        style.configure("Panel.TLabel", background=t["panel"], foreground=t["text"], font=FONT_UI)
        style.configure("Dim.TLabel", background=t["panel"], foreground=t["text_dim"], font=FONT_UI_DIM)
        style.configure("Header.TLabel", background=t["panel"], foreground=t["text_dim"], font=FONT_UI_BOLD_SM)
        style.configure("TButton", background=t["panel2"], foreground=t["text"], borderwidth=1,
                         bordercolor=t["border"], focuscolor=t["accent"], padding=6)
        style.map("TButton", background=[("active", t["border"]), ("disabled", t["panel2"])],
                  foreground=[("disabled", t["text_dim"])])
        style.configure("TCombobox", fieldbackground=t["panel2"], background=t["panel2"],
                         foreground=t["text"], arrowcolor=t["text_dim"])
        style.map("TCombobox", fieldbackground=[("readonly", t["panel2"])],
                  foreground=[("readonly", t["text"])])
        style.configure("TEntry", fieldbackground=t["panel2"], foreground=t["text"],
                         insertcolor=t["text"], bordercolor=t["border"])
        style.map("TEntry", bordercolor=[("focus", t["accent"])])
        style.configure("TScrollbar", background=t["panel2"], troughcolor=t["panel"],
                         bordercolor=t["panel"], arrowcolor=t["text_dim"])
        style.configure("Horizontal.TProgressbar", background=t["accent"], troughcolor=t["panel2"],
                         bordercolor=t["panel"], lightcolor=t["accent"], darkcolor=t["accent"])
        style.configure("TNotebook", background=t["panel"], borderwidth=0)
        style.configure("TNotebook.Tab", background=t["panel2"], foreground=t["text_dim"],
                         padding=(14, 7), borderwidth=0, font=FONT_UI_DIM)
        style.map("TNotebook.Tab", background=[("selected", t["panel"])],
                  foreground=[("selected", t["text"])])

        for widget, role in self._raw_frames:
            try:
                if role == "panel":
                    widget.configure(bg=t["panel"], highlightbackground=t["border"])
                elif role == "panel2":
                    widget.configure(bg=t["panel2"], highlightbackground=t["border"])
                elif role == "bg":
                    widget.configure(bg=t["bg"])
            except tk.TclError:
                pass

        for widget, kind in self._raw_texts:
            try:
                if kind == "text":
                    widget.configure(bg=t["panel2"], fg=t["text"], insertbackground=t["text"],
                                      selectbackground=t["accent"], selectforeground=t["accent_text"])
                elif kind == "listbox":
                    widget.configure(bg=t["panel2"], fg=t["text"], selectbackground=t["accent"],
                                      selectforeground=t["accent_text"])
                elif kind == "entry":
                    widget.configure(bg=t["panel2"], fg=t["text"], insertbackground=t["text"])
            except tk.TclError:
                pass

        for tag_widget in getattr(self, "_tag_targets", []):
            tag_widget.tag_configure("hit", foreground=t["hit"], font=(FONT_MONO[0], FONT_MONO[1], "bold"))
            tag_widget.tag_configure("time", foreground=t["text_dim"])

        if hasattr(self, "menubar"):
            self._theme_menu(self.menubar)

        for w in self._themed_widgets:
            w.redraw()

    def _theme_menu(self, menu):
        t = self.theme
        try:
            menu.configure(bg=t["panel"], fg=t["text"], activebackground=t["accent"],
                            activeforeground=t["accent_text"])
        except tk.TclError:
            pass
        end = menu.index("end")
        if end is None:
            return
        for i in range(end + 1):
            try:
                sub = menu.entrycget(i, "menu")
            except tk.TclError:
                sub = ""
            if sub:
                self._theme_menu(self.nametowidget(sub))

    def _on_theme_choice(self):
        self.settings["theme"] = self.theme_var.get()
        save_settings(self.settings)
        self.theme = THEMES[self._effective_theme_name()]
        self._apply_theme()

    # ---------- меню ----------

    def _build_menu(self):
        menubar = tk.Menu(self, tearoff=False)
        self.menubar = menubar

        file_menu = tk.Menu(menubar, tearoff=False)
        file_menu.add_command(label=t("Выбрать файлы..."), accelerator="Ctrl+O", command=self.pick_files)
        file_menu.add_command(label=t("Выбрать папку..."), accelerator="Ctrl+Shift+O", command=self.pick_folder)
        file_menu.add_separator()
        file_menu.add_command(label=t("Экспорт отчёта (CSV)..."), command=self.export_csv)
        file_menu.add_command(label=t("Экспорт отчёта (Excel)..."), command=self.export_xlsx)
        file_menu.add_command(label=t("Отчёт по всем группам..."), command=self.export_all_groups)
        file_menu.add_command(label=t("Скопировать сводку в буфер"), command=self.copy_summary_to_clipboard)
        file_menu.add_separator()
        file_menu.add_command(label=t("Выход"), accelerator="Ctrl+Q", command=self._on_close)
        menubar.add_cascade(label=t("Файл"), menu=file_menu)

        edit_menu = tk.Menu(menubar, tearoff=False)
        edit_menu.add_command(label=t("Новая группа..."), command=self.add_group)
        edit_menu.add_command(label=t("Переименовать группу..."), command=self.rename_group)
        edit_menu.add_command(label=t("Удалить группу"), command=self.delete_group)
        edit_menu.add_separator()
        edit_menu.add_command(label=t("Сохранить группу..."), accelerator="Ctrl+S", command=self.save_keywords)
        edit_menu.add_command(label=t("Загрузить в группу..."), command=self.load_keywords)
        edit_menu.add_separator()
        self.profiles_menu = tk.Menu(edit_menu, tearoff=False, postcommand=self._build_profiles_menu)
        edit_menu.add_cascade(label=t("Профили"), menu=self.profiles_menu)
        menubar.add_cascade(label=t("Правка"), menu=edit_menu)

        view_menu = tk.Menu(menubar, tearoff=False)
        view_menu.add_command(label=t("Светлая"), command=lambda: self._set_theme("light"))
        view_menu.add_command(label=t("Тёмная"), command=lambda: self._set_theme("dark"))
        view_menu.add_command(label=t("Как в системе"), command=lambda: self._set_theme("system"))
        view_menu.add_separator()
        view_menu.add_command(label=t("Переключить тему"), accelerator="Ctrl+D", command=self.toggle_theme)
        view_menu.add_separator()
        view_menu.add_command(label=("English" if LANG == "ru" else "Русский"), command=self.toggle_lang)
        menubar.add_cascade(label=t("Вид"), menu=view_menu)

        log_menu = tk.Menu(menubar, tearoff=False)
        log_menu.add_command(label=t("Анализ без сохранения"), accelerator="Ctrl+Shift+A", command=self.analyze_only)
        log_menu.add_command(label=t("История обработок..."), command=self.show_history)
        menubar.add_cascade(label=t("Журнал"), menu=log_menu)

        help_menu = tk.Menu(menubar, tearoff=False)
        help_menu.add_command(label=t("Проверить обновления"), command=self.check_updates)
        help_menu.add_command(label=t("О программе"), command=self.show_about)
        menubar.add_cascade(label=t("Справка"), menu=help_menu)

        self.config(menu=menubar)

    def _set_theme(self, name):
        self.theme_var.set(name)
        if hasattr(self, "theme_control"):
            self.theme_control.redraw()
        self._on_theme_choice()

    def toggle_lang(self):
        new_lang = "en" if self.settings.get("lang", "ru") == "ru" else "ru"
        self.settings["lang"] = new_lang
        save_settings(self.settings)
        msg = ("Restart Subtitle Filter to apply the language change."
               if new_lang == "en" else
               "Перезапусти Subtitle Filter, чтобы применить смену языка.")
        messagebox.showinfo(APP_TITLE, msg)

    def toggle_theme(self):
        order = ["light", "dark", "system"]
        current = self.theme_var.get()
        nxt = order[(order.index(current) + 1) % len(order)] if current in order else "light"
        self._set_theme(nxt)

    def _build_profiles_menu(self):
        menu = self.profiles_menu
        menu.delete(0, "end")
        menu.add_command(label=t("Сохранить текущие настройки как профиль..."), command=self.save_profile)
        profiles = self.settings.get("profiles", {})
        if profiles:
            menu.add_separator()
            for name in sorted(profiles.keys()):
                sub = tk.Menu(menu, tearoff=False)
                sub.add_command(label=t("Загрузить"), command=lambda n=name: self.load_profile(n))
                sub.add_command(label=t("Удалить"), command=lambda n=name: self.delete_profile(n))
                menu.add_cascade(label=name, menu=sub)
        self._theme_menu(menu)

    def _bind_shortcuts(self):
        self.bind_all("<Control-o>", lambda e: self.pick_files())
        self.bind_all("<Control-O>", lambda e: self.pick_files())
        self.bind_all("<Control-Shift-o>", lambda e: self.pick_folder())
        self.bind_all("<Control-Shift-O>", lambda e: self.pick_folder())
        self.bind_all("<Control-Return>", lambda e: self.process_files())
        self.bind_all("<Control-s>", lambda e: self.save_keywords())
        self.bind_all("<Control-S>", lambda e: self.save_keywords())
        self.bind_all("<Control-d>", lambda e: self.toggle_theme())
        self.bind_all("<Control-D>", lambda e: self.toggle_theme())
        self.bind_all("<Control-q>", lambda e: self._on_close())
        self.bind_all("<Control-Shift-A>", lambda e: self.analyze_only())
        self.bind_all("<Control-Shift-a>", lambda e: self.analyze_only())
        self.file_listbox.bind("<Delete>", lambda e: self.remove_selected_files())

    def show_about(self):
        win = tk.Toplevel(self)
        win.title(t("О программе"))
        win.configure(bg=self.theme["panel"])
        win.geometry("360x220")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text=APP_TITLE, style="Panel.TLabel", font=("Segoe UI", 14, "bold")).pack(pady=(20, 4))
        ttk.Label(win, text=(t('Версия ') + f"{APP_VERSION}"), style="Panel.TLabel").pack()
        ttk.Label(win, text=t("Оставляет в субтитрах только то, что важно."),
                  style="Dim.TLabel", wraplength=300, justify="center").pack(pady=(10, 10))
        link = tk.Label(win, text=t("Открыть репозиторий на GitHub"), fg=self.theme["accent"],
                         bg=self.theme["panel"], cursor="hand2", font=FONT_UI)
        link.pack()
        link.bind("<Button-1>", lambda e: webbrowser.open(
            f"https://github.com/{GITHUB_OWNER}/{GITHUB_REPO}"))
        ttk.Button(win, text=t("Закрыть"), command=win.destroy).pack(pady=16)

    def _on_close(self):
        try:
            self.settings["window_geometry"] = self.geometry()
            save_settings(self.settings)
        except Exception:
            pass
        self.destroy()

    # ---------- построение интерфейса ----------

    def _mk_border_frame(self, parent, role="panel"):
        f = tk.Frame(parent, highlightthickness=1)
        self._raw_frames.append((f, role))
        return f

    def _mk_panel(self, parent, radius=12):
        p = RoundedPanel(parent, self, radius=radius)
        return p, p.content

    def _build_ui(self):
        pad = 12

        header = ttk.Frame(self, style="TFrame")
        header.pack(fill="x", padx=pad, pady=(pad, 4))
        top_row = ttk.Frame(header, style="TFrame")
        top_row.pack(fill="x")
        ttk.Label(top_row, text=APP_TITLE, font=("Segoe UI", 16, "bold")).pack(side="left")
        ttk.Button(top_row, text=t("Проверить обновления"), command=self.check_updates).pack(side="right")
        self.theme_control = SegmentedControl(
            top_row, self,
            [("light", "☀️"), ("dark", "🌙"), ("system", "💻")],
            self.theme_var, command=self._on_theme_choice, width=110, height=28,
        )
        self.theme_control.pack(side="right", padx=(0, 10))
        dnd_note = t("перетаскивание файлов включено") if DND_AVAILABLE else t("перетаскивание выключено (нет tkinterdnd2)")
        ttk.Label(
            header,
            text=(t('Оставляет в .srt/.vtt/.ass/.ssa только строки с нужными словами (или убирает их). ') + f"{dnd_note}" + '.'),
            foreground=THEMES["light"]["text_dim"],
        ).pack(anchor="w", pady=(2, 0))

        body = ttk.Frame(self, style="TFrame")
        body.pack(fill="both", expand=True, padx=pad, pady=6)
        body.columnconfigure(0, weight=1, minsize=340)
        body.columnconfigure(1, weight=2)
        body.rowconfigure(0, weight=1)

        self._build_left_panel(body)
        self._build_right_panel(body)
        self._build_bottom_panel()

    # ---- левая панель: группы слов, исключения, настройки поиска ----

    def _build_left_panel(self, parent):
        panel, left = self._mk_panel(parent, radius=12)
        panel.grid(row=0, column=0, sticky="nsew", padx=(0, 8))

        notebook = ttk.Notebook(left)
        notebook.pack(fill="both", expand=True)

        words_tab = ttk.Frame(notebook, style="Panel.TFrame")
        settings_tab = ttk.Frame(notebook, style="Panel.TFrame")
        notebook.add(words_tab, text=t("Слова"))
        notebook.add(settings_tab, text=t("Настройки поиска"))

        self._build_words_tab(words_tab)
        self._build_settings_tab(settings_tab)

    def _build_words_tab(self, left):
        ttk.Label(left, text=t("ГРУППЫ СЛОВ"), style="Header.TLabel").pack(anchor="w", padx=4, pady=(10, 4))

        group_row = ttk.Frame(left, style="Panel.TFrame")
        group_row.pack(fill="x", padx=4)
        self.group_combo = ttk.Combobox(group_row, state="readonly", values=[])
        self.group_combo.pack(side="left", fill="x", expand=True)
        self.group_combo.bind("<<ComboboxSelected>>", lambda e: self._select_group(self.group_combo.get()))

        group_btns = ttk.Frame(left, style="Panel.TFrame")
        group_btns.pack(fill="x", padx=4, pady=(6, 0))
        ttk.Button(group_btns, text=t("+ группа"), command=self.add_group).pack(side="left", padx=(0, 4))
        ttk.Button(group_btns, text=t("Переименовать"), command=self.rename_group).pack(side="left", padx=(0, 4))
        ttk.Button(group_btns, text=t("Удалить"), command=self.delete_group).pack(side="left")

        ttk.Label(
            left,
            text=('«' + f"{ALL_GROUPS_LABEL}" + t('» объединяет слова из всех групп сразу.')),
            style="Dim.TLabel", wraplength=280, justify="left",
        ).pack(anchor="w", padx=4, pady=(6, 6))

        ttk.Label(left, text=t("СЛОВА В ГРУППЕ (по одному в строке)"), style="Header.TLabel").pack(
            anchor="w", padx=4, pady=(4, 4)
        )
        self.kw_text = tk.Text(left, relief="flat", font=FONT_MONO_KW, wrap="word", height=7, undo=True, maxundo=100)
        self._raw_texts.append((self.kw_text, "text"))
        self.kw_text.pack(fill="both", expand=True, padx=4)
        self.kw_text.bind("<<Modified>>", self._on_kw_change)

        self.kw_count_lbl = ttk.Label(left, text=t("0 слов"), style="Dim.TLabel")
        self.kw_count_lbl.pack(anchor="w", padx=4, pady=(4, 8))

        btn_row = ttk.Frame(left, style="Panel.TFrame")
        btn_row.pack(fill="x", padx=4, pady=(0, 10))
        ttk.Button(btn_row, text=t("Загрузить..."), command=self.load_keywords).pack(side="left", padx=(0, 6))
        ttk.Button(btn_row, text=t("Сохранить..."), command=self.save_keywords).pack(side="left")

        ttk.Label(left, text=t("ИСКЛЮЧИТЬ СЛОВА (всегда вырезаются из результата)"), style="Header.TLabel").pack(
            anchor="w", padx=4, pady=(2, 4)
        )
        self.exclude_text = tk.Text(left, relief="flat", font=FONT_MONO_KW, wrap="word", height=4, undo=True, maxundo=100)
        self._raw_texts.append((self.exclude_text, "text"))
        self.exclude_text.pack(fill="x", padx=4, pady=(0, 10))
        self.exclude_text.bind("<<Modified>>", self._on_exclude_change)

    def _build_settings_tab(self, left):
        ttk.Label(left, text=t("НАСТРОЙКИ ПОИСКА"), style="Header.TLabel").pack(anchor="w", padx=4, pady=(10, 6))

        self._build_toggle_row(left, t("Только целые слова"), self.whole_word_var, self._on_whole_word_change)
        self._build_toggle_row(left, t("Учитывать регистр букв"), self.case_sensitive_var, self._on_search_opt_change)
        self._build_toggle_row(left, t("Регулярные выражения (regex)"), self.regex_mode_var, self._on_regex_toggle)

        mode_row = ttk.Frame(left, style="Panel.TFrame")
        mode_row.pack(fill="x", padx=4, pady=(10, 4))
        ttk.Label(mode_row, text=t("Режим:"), style="Panel.TLabel").pack(side="left", padx=(0, 8))
        self.mode_control = SegmentedControl(
            mode_row, self, [(MODE_KEEP, t("Оставить")), (MODE_REMOVE, t("Убрать"))],
            self.mode_var, command=self._on_mode_change, width=180, height=26,
        )
        self.mode_control.pack(side="left")

        self.mode_hint_lbl = ttk.Label(left, text="", style="Dim.TLabel", wraplength=280, justify="left")
        self.mode_hint_lbl.pack(anchor="w", padx=4, pady=(6, 12))

        ttk.Label(left, text=t("ФИЛЬТР ПО ВРЕМЕНИ"), style="Header.TLabel").pack(anchor="w", padx=4, pady=(0, 6))
        self._build_toggle_row(left, t("Ограничить диапазон времени"), self.time_filter_var, self._on_time_filter_toggle)

        time_row = ttk.Frame(left, style="Panel.TFrame")
        time_row.pack(fill="x", padx=4, pady=(4, 4))
        ttk.Label(time_row, text=t("От"), style="Panel.TLabel").pack(side="left")
        self.time_start_entry = ttk.Entry(time_row, textvariable=self.time_start_var, width=9)
        self.time_start_entry.pack(side="left", padx=(4, 10))
        ttk.Label(time_row, text=t("До"), style="Panel.TLabel").pack(side="left")
        self.time_end_entry = ttk.Entry(time_row, textvariable=self.time_end_var, width=9)
        self.time_end_entry.pack(side="left", padx=(4, 0))
        self.time_start_var.trace_add("write", lambda *a: self._on_time_bounds_change())
        self.time_end_var.trace_add("write", lambda *a: self._on_time_bounds_change())

        ttk.Label(left, text=t("Формат: ЧЧ:ММ:СС, например 00:05:00"), style="Dim.TLabel").pack(
            anchor="w", padx=4, pady=(2, 12)
        )

    def _build_toggle_row(self, parent, label, var, command):
        row = ttk.Frame(parent, style="Panel.TFrame")
        row.pack(fill="x", padx=4, pady=3)
        ttk.Label(row, text=label, style="Panel.TLabel").pack(side="left")
        toggle = ToggleSwitch(row, self, var, command=command)
        toggle.pack(side="right")
        return toggle

    # ---- правая панель: файлы ----

    def _build_right_panel(self, parent):
        panel, right = self._mk_panel(parent, radius=12)
        panel.grid(row=0, column=1, sticky="nsew")
        right.rowconfigure(3, weight=1)
        right.columnconfigure(0, weight=1)

        ttk.Label(right, text=t("ФАЙЛЫ СУБТИТРОВ (.srt / .vtt / .ass / .ssa)"), style="Header.TLabel").grid(
            row=0, column=0, sticky="w", padx=4, pady=(10, 6)
        )

        file_btns = ttk.Frame(right, style="Panel.TFrame")
        file_btns.grid(row=1, column=0, sticky="ew", padx=4)
        ttk.Button(file_btns, text=t("Выбрать файлы..."), command=self.pick_files).pack(side="left", padx=(0, 6))
        ttk.Button(file_btns, text=t("Выбрать папку"), command=self.pick_folder).pack(side="left", padx=(0, 6))
        ttk.Button(file_btns, text=t("Удалить выбранное"), command=self.remove_selected_files).pack(side="left", padx=(0, 6))
        ttk.Button(file_btns, text=t("Очистить всё"), command=self.clear_files).pack(side="left")

        search_row = ttk.Frame(right, style="Panel.TFrame")
        search_row.grid(row=2, column=0, sticky="ew", padx=4, pady=(8, 0))
        ttk.Label(search_row, text="🔍", style="Panel.TLabel").pack(side="left")
        search_entry = ttk.Entry(search_row, textvariable=self.search_var)
        search_entry.pack(side="left", fill="x", expand=True, padx=(4, 0))

        list_frame = self._mk_border_frame(right, "panel2")
        list_frame.grid(row=3, column=0, sticky="nsew", padx=4, pady=10)
        list_frame.rowconfigure(0, weight=1)
        list_frame.columnconfigure(0, weight=1)

        self.file_listbox = tk.Listbox(
            list_frame, relief="flat", font=FONT_MONO, activestyle="none", selectmode="extended",
        )
        self._raw_texts.append((self.file_listbox, "listbox"))
        self.file_listbox.grid(row=0, column=0, sticky="nsew")
        self.file_listbox.bind("<<ListboxSelect>>", self._on_file_select)
        self.file_listbox.bind("<Button-3>", self._show_file_context_menu)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.file_listbox.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.file_listbox.config(yscrollcommand=scrollbar.set)

        self.file_count_lbl = ttk.Label(right, text=t("Файлов не выбрано (можно перетащить сюда)"), style="Dim.TLabel")
        self.file_count_lbl.grid(row=4, column=0, sticky="w", padx=4)

        action_row = ttk.Frame(right, style="Panel.TFrame")
        action_row.grid(row=5, column=0, sticky="ew", padx=4, pady=(8, 6))
        self.process_btn = PillButton(action_row, self, t("Обработать"), self.process_files)
        self.process_btn.pack(side="left")
        ttk.Button(action_row, text=t("Анализ без сохранения"), command=self.analyze_only).pack(side="left", padx=(10, 0))
        self.open_folder_btn = ttk.Button(
            action_row, text=t("Открыть папку с результатом"), command=self.open_output_folder, state="disabled"
        )
        self.open_folder_btn.pack(side="left", padx=(8, 0))

        action_row2 = ttk.Frame(right, style="Panel.TFrame")
        action_row2.grid(row=6, column=0, sticky="ew", padx=4, pady=(0, 6))
        self.export_csv_btn = ttk.Button(action_row2, text="CSV", command=self.export_csv, state="disabled")
        self.export_csv_btn.pack(side="left")
        self.export_xlsx_btn = ttk.Button(action_row2, text="Excel", command=self.export_xlsx, state="disabled")
        self.export_xlsx_btn.pack(side="left", padx=(6, 0))
        self.copy_btn = ttk.Button(action_row2, text=t("Скопировать сводку"), command=self.copy_summary_to_clipboard, state="disabled")
        self.copy_btn.pack(side="left", padx=(6, 0))

        progress_row = ttk.Frame(right, style="Panel.TFrame")
        progress_row.grid(row=7, column=0, sticky="ew", padx=4, pady=(0, 10))
        self.progress = ttk.Progressbar(progress_row, orient="horizontal", mode="determinate")
        self.progress.pack(fill="x")

    # ---- нижняя панель: лог + предпросмотр ----

    def _build_bottom_panel(self):
        bottom = ttk.Frame(self, style="TFrame")
        bottom.pack(fill="both", expand=False, padx=12, pady=(0, 12))
        bottom.columnconfigure(0, weight=1)
        bottom.columnconfigure(1, weight=1)

        log_panel, log_frame = self._mk_panel(bottom, radius=12)
        log_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        ttk.Label(log_frame, text=t("РЕЗУЛЬТАТ ОБРАБОТКИ"), style="Header.TLabel").pack(anchor="w", padx=2, pady=(8, 4))
        self.log_text = tk.Text(log_frame, relief="flat", font=FONT_MONO, height=9, state="disabled")
        self._raw_texts.append((self.log_text, "text"))
        self.log_text.pack(fill="both", expand=True, padx=2, pady=(0, 8))

        preview_panel, preview_frame = self._mk_panel(bottom, radius=12)
        preview_panel.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
        ttk.Label(
            preview_frame,
            text=t("ПРЕДПРОСМОТР (двойной клик по строке — открыть видео рядом на этом месте)"),
            style="Header.TLabel",
        ).pack(anchor="w", padx=2, pady=(8, 4))
        self.preview_text = tk.Text(preview_frame, relief="flat", font=FONT_MONO, height=9, state="disabled", cursor="hand2")
        self._raw_texts.append((self.preview_text, "text"))
        self.preview_text.pack(fill="both", expand=True, padx=2, pady=(0, 8))
        self.preview_text.bind("<Double-Button-1>", self._on_preview_double_click)
        self._tag_targets = [self.preview_text]
        self.preview_text.tag_configure("hit")
        self.preview_text.tag_configure("time")

    # ---------- группы слов ----------

    def _load_group_names(self):
        names = list(self.settings["groups"].keys())
        self.group_combo["values"] = [ALL_GROUPS_LABEL] + names

    def _select_group(self, name):
        if not name:
            name = DEFAULT_GROUP
        self.group_combo.set(name)
        self.settings["active_group"] = name
        self.kw_text.config(state="normal")
        self.kw_text.delete("1.0", "end")
        if name == ALL_GROUPS_LABEL:
            words = self._all_group_words()
            self.kw_text.insert("1.0", "\n".join(words))
            self.kw_text.config(state="disabled")
        else:
            words = self.settings["groups"].get(name, [])
            self.kw_text.insert("1.0", "\n".join(words))
        self._on_kw_change(save=False)

    def _all_group_words(self):
        seen = []
        for words in self.settings["groups"].values():
            for w in words:
                if w not in seen:
                    seen.append(w)
        return seen

    def add_group(self):
        name = self._ask_text(t("Новая группа"), t("Название новой группы слов:"))
        if not name:
            return
        if name in self.settings["groups"] or name == ALL_GROUPS_LABEL:
            messagebox.showwarning(APP_TITLE, t("Группа с таким названием уже есть."))
            return
        self.settings["groups"][name] = []
        self._load_group_names()
        self._select_group(name)
        save_settings(self.settings)

    def rename_group(self):
        current = self.group_combo.get()
        if current == ALL_GROUPS_LABEL:
            messagebox.showinfo(APP_TITLE, t("«Все группы» — это не отдельная группа, переименовать нельзя."))
            return
        new_name = self._ask_text(t("Переименовать группу"), t("Новое название:"), initial=current)
        if not new_name or new_name == current:
            return
        if new_name in self.settings["groups"]:
            messagebox.showwarning(APP_TITLE, t("Группа с таким названием уже есть."))
            return
        self.settings["groups"][new_name] = self.settings["groups"].pop(current)
        for path, g in list(self.settings.get("file_groups", {}).items()):
            if g == current:
                self.settings["file_groups"][path] = new_name
        self._load_group_names()
        self._select_group(new_name)
        save_settings(self.settings)

    def delete_group(self):
        current = self.group_combo.get()
        if current == ALL_GROUPS_LABEL:
            return
        if len(self.settings["groups"]) <= 1:
            messagebox.showwarning(APP_TITLE, t("Нельзя удалить единственную группу."))
            return
        if not messagebox.askyesno(APP_TITLE, (t('Удалить группу «') + f"{current}" + '»?')):
            return
        del self.settings["groups"][current]
        self.settings["file_groups"] = {
            p: g for p, g in self.settings.get("file_groups", {}).items() if g != current
        }
        self._load_group_names()
        self._select_group(next(iter(self.settings["groups"])))
        save_settings(self.settings)
        self._refresh_file_list()

    def _ask_text(self, title, prompt, initial=""):
        t = self.theme
        win = tk.Toplevel(self)
        win.title(title)
        win.configure(bg=t["panel"])
        win.geometry("320x120")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text=prompt, style="Panel.TLabel").pack(padx=12, pady=(14, 6), anchor="w")
        entry = tk.Entry(win, bg=t["panel2"], fg=t["text"], insertbackground=t["text"], relief="flat")
        entry.insert(0, initial)
        entry.pack(fill="x", padx=12)
        entry.focus_set()
        result = {"value": None}

        def confirm(event=None):
            result["value"] = entry.get().strip()
            win.destroy()

        entry.bind("<Return>", confirm)
        btns = ttk.Frame(win, style="Panel.TFrame")
        btns.pack(pady=12)
        ttk.Button(btns, text=t("ОК"), command=confirm).pack(side="left", padx=4)
        ttk.Button(btns, text=t("Отмена"), command=win.destroy).pack(side="left", padx=4)
        self.wait_window(win)
        return result["value"]

    def _ask_choice(self, title, prompt, options, initial=None):
        t = self.theme
        win = tk.Toplevel(self)
        win.title(title)
        win.configure(bg=t["panel"])
        win.geometry("320x140")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text=prompt, style="Panel.TLabel").pack(padx=12, pady=(14, 6), anchor="w")
        combo = ttk.Combobox(win, state="readonly", values=options)
        combo.set(initial if initial in options else (options[0] if options else ""))
        combo.pack(fill="x", padx=12)
        result = {"value": None}

        def confirm():
            result["value"] = combo.get()
            win.destroy()

        btns = ttk.Frame(win, style="Panel.TFrame")
        btns.pack(pady=14)
        ttk.Button(btns, text=t("ОК"), command=confirm).pack(side="left", padx=4)
        ttk.Button(btns, text=t("Отмена"), command=win.destroy).pack(side="left", padx=4)
        self.wait_window(win)
        return result["value"]

    # ---------- обработчики настроек поиска ----------

    def _on_kw_change(self, event=None, save=True):
        self.kw_text.edit_modified(False)
        words = self._get_keywords()
        self.kw_count_lbl.config(text=(f"{len(words)}" + t(' слов(а)')))
        current = self.group_combo.get()
        if save and current != ALL_GROUPS_LABEL:
            self.settings["groups"][current] = words
            save_settings(self.settings)
        self._schedule_preview()

    def _load_exclude_words(self):
        self.exclude_text.delete("1.0", "end")
        self.exclude_text.insert("1.0", "\n".join(self.settings.get("exclude_words", [])))
        self.exclude_text.edit_modified(False)

    def _on_exclude_change(self, event=None):
        self.exclude_text.edit_modified(False)
        words = [w.strip() for w in self.exclude_text.get("1.0", "end").splitlines() if w.strip()]
        self.settings["exclude_words"] = words
        save_settings(self.settings)
        self._schedule_preview()

    def _on_whole_word_change(self):
        self.settings["whole_word"] = self.whole_word_var.get()
        save_settings(self.settings)
        self._schedule_preview()

    def _on_search_opt_change(self):
        self.settings["case_sensitive"] = self.case_sensitive_var.get()
        save_settings(self.settings)
        self._schedule_preview()

    def _on_regex_toggle(self):
        self.settings["regex_mode"] = self.regex_mode_var.get()
        save_settings(self.settings)
        self._sync_control_states()
        self._schedule_preview()

    def _on_mode_change(self):
        self.settings["mode"] = self.mode_var.get()
        save_settings(self.settings)
        self._sync_control_states()
        self._schedule_preview()

    def _on_time_filter_toggle(self):
        self.settings["time_filter_enabled"] = self.time_filter_var.get()
        save_settings(self.settings)
        self._schedule_preview()

    def _on_time_bounds_change(self):
        self.settings["time_filter_start"] = self.time_start_var.get()
        self.settings["time_filter_end"] = self.time_end_var.get()
        save_settings(self.settings)
        self._schedule_preview()

    def _time_bounds(self):
        start = time_to_seconds(self.time_start_var.get()) if self.time_start_var.get().strip() else None
        end = time_to_seconds(self.time_end_var.get()) if self.time_end_var.get().strip() else None
        return self.time_filter_var.get(), start, end

    def _sync_control_states(self):
        enabled = not self.regex_mode_var.get()
        for w in self._themed_widgets:
            if isinstance(w, ToggleSwitch) and w.var is self.whole_word_var:
                w.set_enabled(enabled)
        if self.mode_var.get() == MODE_REMOVE:
            self.mode_hint_lbl.config(
                text=t("«Убрать» — из субтитров удаляются строки с совпадениями, всё остальное остаётся.")
            )
        else:
            self.mode_hint_lbl.config(
                text=t("«Оставить» — в субтитрах остаются только строки с совпадениями (обычный режим).")
            )

    def _get_keywords(self):
        raw = self.kw_text.get("1.0", "end")
        return [w.strip() for w in raw.splitlines() if w.strip()]

    def _active_keywords(self):
        if self.group_combo.get() == ALL_GROUPS_LABEL:
            return self._all_group_words()
        return self._get_keywords()

    def _keywords_for_file(self, path):
        t("""Учитывает персональную группу файла (если назначена через контекстное меню).""")
        override = self.settings.get("file_groups", {}).get(path)
        if override and override in self.settings["groups"]:
            return self.settings["groups"][override]
        return self._active_keywords()

    def _exclude_words(self):
        return self.settings.get("exclude_words", [])

    def _build_matcher(self, keywords):
        return build_matcher(
            keywords, self.whole_word_var.get(),
            case_sensitive=self.case_sensitive_var.get(),
            regex_mode=self.regex_mode_var.get(),
        )

    def _make_decider(self, keywords):
        matcher = self._build_matcher(keywords)
        exclude_matcher = self._build_matcher(self._exclude_words())
        invert = self.mode_var.get() == MODE_REMOVE
        enabled, t_start, t_end = self._time_bounds()

        def decide(text, start_sec):
            return decide_keep(text, start_sec, matcher, invert, exclude_matcher, enabled, t_start, t_end)

        return decide, matcher, exclude_matcher, invert

    def load_keywords(self):
        if self.group_combo.get() == ALL_GROUPS_LABEL:
            messagebox.showinfo(APP_TITLE, t("Выбери конкретную группу, чтобы загрузить в неё список."))
            return
        path = filedialog.askopenfilename(title=t("Загрузить список слов"), filetypes=[(t("Текстовый файл"), "*.txt")])
        if not path:
            return
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            content = f.read()
        self.kw_text.delete("1.0", "end")
        self.kw_text.insert("1.0", content.strip())
        self._on_kw_change()

    def save_keywords(self):
        path = filedialog.asksaveasfilename(
            title=t("Сохранить список слов"), defaultextension=".txt",
            filetypes=[(t("Текстовый файл"), "*.txt")], initialfile="keywords.txt",
        )
        if not path:
            return
        with open(path, "w", encoding="utf-8") as f:
            f.write(self.kw_text.get("1.0", "end").strip())

    # ---------- профили ----------

    def _profile_snapshot(self):
        return {
            "groups": json.loads(json.dumps(self.settings["groups"])),
            "active_group": self.group_combo.get(),
            "exclude_words": list(self.settings.get("exclude_words", [])),
            "whole_word": self.whole_word_var.get(),
            "case_sensitive": self.case_sensitive_var.get(),
            "regex_mode": self.regex_mode_var.get(),
            "mode": self.mode_var.get(),
            "time_filter_enabled": self.time_filter_var.get(),
            "time_filter_start": self.time_start_var.get(),
            "time_filter_end": self.time_end_var.get(),
        }

    def save_profile(self):
        name = self._ask_text(t("Новый профиль"), t("Название профиля:"))
        if not name:
            return
        self.settings.setdefault("profiles", {})[name] = self._profile_snapshot()
        save_settings(self.settings)
        self.show_toast((t('Профиль «') + f"{name}" + t('» сохранён.')))

    def load_profile(self, name):
        profile = self.settings.get("profiles", {}).get(name)
        if not profile:
            return
        self.settings["groups"] = json.loads(json.dumps(profile["groups"]))
        self.settings["exclude_words"] = list(profile.get("exclude_words", []))
        self.whole_word_var.set(profile.get("whole_word", True))
        self.case_sensitive_var.set(profile.get("case_sensitive", False))
        self.regex_mode_var.set(profile.get("regex_mode", False))
        self.mode_var.set(profile.get("mode", MODE_KEEP))
        self.time_filter_var.set(profile.get("time_filter_enabled", False))
        self.time_start_var.set(profile.get("time_filter_start", ""))
        self.time_end_var.set(profile.get("time_filter_end", ""))
        self._load_group_names()
        self._select_group(profile.get("active_group", DEFAULT_GROUP))
        self._load_exclude_words()
        for w in self._themed_widgets:
            w.redraw()
        self._sync_control_states()
        save_settings(self.settings)
        self.show_toast((t('Профиль «') + f"{name}" + t('» загружен.')))

    def delete_profile(self, name):
        if messagebox.askyesno(APP_TITLE, (t('Удалить профиль «') + f"{name}" + '»?')):
            self.settings.get("profiles", {}).pop(name, None)
            save_settings(self.settings)

    # ---------- файлы ----------

    def pick_files(self):
        paths = filedialog.askopenfilenames(
            title=t("Выбери файлы субтитров"),
            filetypes=[(t("Субтитры"), "*.srt *.vtt *.ass *.ssa")],
        )
        if paths:
            self._add_paths(paths)

    def pick_folder(self):
        folder = filedialog.askdirectory(title=t("Выбери папку с субтитрами"))
        if folder:
            found = []
            for ext in SUPPORTED_EXTS:
                found.extend(glob.glob(os.path.join(folder, f"*{ext}")))
            self._add_paths(sorted(found))
            self.settings["last_folder"] = folder
            save_settings(self.settings)

    def _on_drop_files(self, event):
        paths = self.tk.splitlist(event.data)
        good = [p for p in paths if p.lower().endswith(SUPPORTED_EXTS)]
        self._add_paths(good)

    def _add_paths(self, paths):
        for p in paths:
            if p not in self.files:
                self.files.append(p)
        self._refresh_file_list()

    def clear_files(self):
        self.files = []
        self.search_var.set("")
        self._refresh_file_list()
        self._set_log("")
        self._set_preview("")

    def remove_selected_files(self):
        sel = self.file_listbox.curselection()
        if not sel:
            return
        to_remove = {self._visible_indices[i] for i in sel if i < len(self._visible_indices)}
        self.files = [p for i, p in enumerate(self.files) if i not in to_remove]
        self._refresh_file_list()

    def _show_file_context_menu(self, event):
        idx = self.file_listbox.nearest(event.y)
        if idx < 0 or idx >= len(self._visible_indices):
            return
        if idx not in self.file_listbox.curselection():
            self.file_listbox.selection_clear(0, "end")
            self.file_listbox.selection_set(idx)
        menu = tk.Menu(self, tearoff=False)
        self._theme_menu(menu)
        menu.add_command(label=t("Назначить группу для выбранных..."), command=self._assign_group_to_selected)
        menu.add_command(label=t("Сбросить группу для выбранных"), command=self._clear_group_for_selected)
        menu.add_separator()
        menu.add_command(label=t("Удалить из списка"), command=self.remove_selected_files)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _assign_group_to_selected(self):
        sel = self.file_listbox.curselection()
        if not sel:
            return
        options = [NO_GROUP_OVERRIDE] + list(self.settings["groups"].keys())
        choice = self._ask_choice(t("Группа для файлов"), t("Использовать группу:"), options)
        if not choice:
            return
        file_groups = self.settings.setdefault("file_groups", {})
        for i in sel:
            if i >= len(self._visible_indices):
                continue
            path = self.files[self._visible_indices[i]]
            if choice == NO_GROUP_OVERRIDE:
                file_groups.pop(path, None)
            else:
                file_groups[path] = choice
        save_settings(self.settings)
        self._refresh_file_list()

    def _clear_group_for_selected(self):
        sel = self.file_listbox.curselection()
        file_groups = self.settings.setdefault("file_groups", {})
        for i in sel:
            if i >= len(self._visible_indices):
                continue
            file_groups.pop(self.files[self._visible_indices[i]], None)
        save_settings(self.settings)
        self._refresh_file_list()

    def _refresh_file_list(self):
        query = self.search_var.get().strip().lower()
        file_groups = self.settings.get("file_groups", {})
        self.file_listbox.delete(0, "end")
        self._visible_indices = []
        for i, p in enumerate(self.files):
            name = os.path.basename(p)
            if query and query not in name.lower():
                continue
            self._visible_indices.append(i)
            override = file_groups.get(p)
            label = f"{name}   →  {override}" if override else name
            self.file_listbox.insert("end", label)
        if not self.files:
            text = t("Файлов не выбрано (можно перетащить сюда)")
        elif query and not self._visible_indices:
            text = (t('Ничего не найдено по «') + f"{query}" + '» (' + f"{len(self.files)}" + t(' файл(ов) всего)'))
        else:
            shown = len(self._visible_indices)
            text = (f"{shown}" + t(' из ') + f"{len(self.files)}" + t(' файл(ов)')) if query else (f"{len(self.files)}" + t(' файл(ов) выбрано'))
        self.file_count_lbl.config(text=text)

    def _on_file_select(self, event=None):
        self._schedule_preview()

    # ---------- предпросмотр ----------

    def _schedule_preview(self):
        if self._preview_job:
            self.after_cancel(self._preview_job)
        self._preview_job = self.after(200, self._update_preview)

    def _current_preview_path(self):
        sel = self.file_listbox.curselection()
        if not sel or not self.files or not self._visible_indices:
            return None
        return self.files[self._visible_indices[sel[0]]]

    def _update_preview(self):
        self._preview_job = None
        path = self._current_preview_path()
        if not path:
            self._set_preview(t("Выбери файл в списке слева."))
            return
        keywords = self._keywords_for_file(path)
        if not keywords:
            self._set_preview(t("Список слов пуст."))
            return
        try:
            decide, matcher, exclude_matcher, invert = self._make_decider(keywords)
        except re.error as e:
            self._set_preview((t('Неверное регулярное выражение: ') + f"{e}"))
            return
        try:
            _, entries = peek_subtitle_blocks(path)
        except Exception as e:
            self._set_preview((t('Не удалось прочитать файл: ') + f"{e}"))
            return

        kept = [e for e in entries if decide(e["text"], e["start_sec"])]
        self._render_preview(kept, keywords, highlight=not invert)

    def _render_preview(self, kept, keywords, highlight=True):
        self.preview_text.config(state="normal")
        self.preview_text.delete("1.0", "end")
        if not kept:
            self.preview_text.insert("1.0", t("Совпадений не найдено."))
            self.preview_text.config(state="disabled")
            return

        single_matchers = []
        if highlight:
            for kw in keywords:
                try:
                    single_matchers.append((kw, build_single_matcher(
                        kw, self.whole_word_var.get(),
                        case_sensitive=self.case_sensitive_var.get(),
                        regex_mode=self.regex_mode_var.get(),
                    )))
                except re.error:
                    continue

        for e in kept[:200]:  # ограничение, чтобы не тормозило на огромных файлах
            secs = e["start_sec"]
            time_str = "--:--:--"
            if secs is not None:
                time_str = str(datetime.timedelta(seconds=int(secs)))
                if len(time_str) == 7:
                    time_str = "0" + time_str
            self.preview_text.insert("end", time_str + "\n", "time")
            clean = e["text"]
            pos = 0
            spans = []
            for _, m in single_matchers:
                for match in m.finditer(clean):
                    spans.append((match.start(), match.end()))
            spans.sort()
            merged = []
            for s, en in spans:
                if merged and s <= merged[-1][1]:
                    merged[-1] = (merged[-1][0], max(merged[-1][1], en))
                else:
                    merged.append((s, en))
            for s, en in merged:
                self.preview_text.insert("end", clean[pos:s])
                self.preview_text.insert("end", clean[s:en], "hit")
                pos = en
            self.preview_text.insert("end", clean[pos:] + "\n\n")

        if len(kept) > 200:
            self.preview_text.insert("end", (t('... и ещё ') + f"{len(kept) - 200}" + t(' строк (показаны первые 200)')))
        self.preview_text.config(state="disabled")

    def _set_preview(self, text):
        self.preview_text.config(state="normal")
        self.preview_text.delete("1.0", "end")
        self.preview_text.insert("1.0", text)
        self.preview_text.config(state="disabled")

    def _on_preview_double_click(self, event):
        idx = self.preview_text.index(f"@{event.x},{event.y}")
        click_line = int(idx.split(".")[0])
        ranges = self.preview_text.tag_ranges("time")
        target_start = None
        for i in range(0, len(ranges), 2):
            start = ranges[i]
            if int(str(start).split(".")[0]) <= click_line:
                target_start = start
        if target_start is None:
            return
        end = self.preview_text.index(f"{target_start} lineend")
        time_str = self.preview_text.get(target_start, end).strip()
        seconds = time_to_seconds(time_str)
        if seconds is None:
            return
        self._open_media_at(seconds)

    def _find_vlc(self):
        candidates = [
            r"C:\Program Files\VideoLAN\VLC\vlc.exe",
            r"C:\Program Files (x86)\VideoLAN\VLC\vlc.exe",
            shutil.which("vlc"),
        ]
        for c in candidates:
            if c and os.path.exists(c):
                return c
        return None

    def _open_media_at(self, seconds):
        path = self._current_preview_path()
        if not path:
            return
        base = os.path.splitext(path)[0]
        video = None
        for ext in (".mp4", ".mkv", ".avi", ".mov", ".webm", ".m4v"):
            candidate = base + ext
            if os.path.exists(candidate):
                video = candidate
                break
        if not video:
            self.show_toast(t("Видео с таким же именем рядом с субтитрами не найдено."), kind="warn")
            return
        vlc_path = self._find_vlc()
        try:
            if vlc_path:
                subprocess.Popen([vlc_path, video, f"--start-time={int(seconds)}"])
            else:
                os.startfile(video)
                self.show_toast(t("VLC не найден — видео открыто с начала (перемотка недоступна)."))
        except Exception as e:
            self.show_toast((t('Не удалось открыть видео: ') + f"{e}"), kind="error")

    # ---------- лог ----------

    def _set_log(self, text):
        self.log_text.config(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.insert("1.0", text)
        self.log_text.config(state="disabled")

    def _append_log(self, line):
        self.log_text.config(state="normal")
        self.log_text.insert("end", line + "\n")
        self.log_text.see("end")
        self.log_text.config(state="disabled")

    # ---------- анализ без сохранения ----------

    def analyze_only(self):
        if not self.files:
            messagebox.showwarning(APP_TITLE, t("Сначала выбери файлы субтитров."))
            return
        self._set_log(t("Анализ...\n"))
        totals = {}
        total_kept = total_all = 0
        errors = 0
        for path in self.files:
            keywords = self._keywords_for_file(path)
            if not keywords:
                continue
            try:
                decide, matcher, exclude_matcher, invert = self._make_decider(keywords)
                single_matchers = [
                    (kw, build_single_matcher(kw, self.whole_word_var.get(),
                                               self.case_sensitive_var.get(), self.regex_mode_var.get()))
                    for kw in keywords
                ]
                _, entries = peek_subtitle_blocks(path)
            except (re.error, Exception) as e:
                self._append_log(('✕ ' + f"{os.path.basename(path)}" + t(' — ошибка: ') + f"{e}"))
                errors += 1
                continue
            kept = [e for e in entries if decide(e["text"], e["start_sec"])]
            all_text = "\n".join(e["text"] for e in entries)
            for kw, m in single_matchers:
                c = len(m.findall(all_text))
                if c:
                    totals[kw] = totals.get(kw, 0) + c
            total_kept += len(kept)
            total_all += len(entries)
            self._append_log(('• ' + f"{os.path.basename(path)}" + t(' — совпало бы ') + f"{len(kept)}" + t(' из ') + f"{len(entries)}" + t(' строк')))

        self._append_log((t('\nИтого без сохранения: ') + f"{total_kept}" + t(' из ') + f"{total_all}" + t(' строк.')))
        if totals:
            top = sorted(totals.items(), key=lambda kv: kv[1], reverse=True)[:8]
            self._append_log(t("Чаще всего встречалось: ") + ", ".join(f"«{k}» — {v}" for k, v in top))
        self.show_toast(t("Анализ завершён — файлы не изменены."))

    # ---------- обработка ----------

    def process_files(self):
        if not self.files:
            messagebox.showwarning(APP_TITLE, t("Сначала выбери файлы субтитров."))
            return
        if not any(self._keywords_for_file(p) for p in self.files):
            messagebox.showwarning(APP_TITLE, t("Сначала добавь хотя бы одно слово в список."))
            return

        invert = self.mode_var.get() == MODE_REMOVE
        suffix = "cleaned" if invert else "filtered"
        first_dir = os.path.dirname(self.files[0])
        out_dir = os.path.join(first_dir, "filtered")
        os.makedirs(out_dir, exist_ok=True)

        self.process_btn.set_enabled(False)
        self._set_log(t("Обработка...\n"))
        self.progress.config(maximum=len(self.files), value=0)
        total_files = len(self.files)

        def worker():
            report_totals = {}
            total_kept = 0
            total_blocks = 0
            failures = 0

            for idx, path in enumerate(self.files, start=1):
                keywords = self._keywords_for_file(path)
                try:
                    if not keywords:
                        raise ValueError(t("для этого файла список слов пуст"))
                    decide, matcher, exclude_matcher, file_invert = self._make_decider(keywords)
                    single_matchers = [
                        (kw, build_single_matcher(kw, self.whole_word_var.get(),
                                                   self.case_sensitive_var.get(), self.regex_mode_var.get()))
                        for kw in keywords
                    ]
                    out_path, kept, total, plain_text = filter_subtitle_file(path, decide, out_dir, suffix)
                    for kw, m in single_matchers:
                        count = len(m.findall(plain_text))
                        if count > 0:
                            self.last_report.append({"file": os.path.basename(path), "keyword": kw, "count": count})
                            report_totals[kw] = report_totals.get(kw, 0) + count
                    name = os.path.basename(out_path)
                    line = ('✓ ' + f"{name}" + ' — ' + f"{kept}" + t(' из ') + f"{total}" + t(' строк')) if kept else \
                           ('⚠ ' + f"{name}" + t(' — 0 из ') + f"{total}" + t(' строк (совпадений нет)'))
                    total_kept += kept
                    total_blocks += total
                except Exception as e:
                    line = ('✕ ' + f"{os.path.basename(path)}" + t(' — ошибка: ') + f"{e}")
                    failures += 1

                self.after(0, lambda i=idx, ln=line: (self._append_log(ln), self.progress.config(value=i)))

            def finish():
                self._append_log((t('\nИтого: ') + f"{total_kept}" + t(' из ') + f"{total_blocks}" + t(' строк в ') + f"{total_files - failures}" + t(' файл(ах).')))
                if report_totals:
                    top = sorted(report_totals.items(), key=lambda kv: kv[1], reverse=True)[:5]
                    top_str = ", ".join(f"«{k}» — {v}" for k, v in top)
                    self._append_log((t('Чаще всего встречалось: ') + f"{top_str}"))
                self._append_log((t('\nГотово. Результаты в папке:\n') + f"{out_dir}"))
                self.last_out_dir = out_dir
                self.open_folder_btn.config(state="normal")
                has_report = bool(self.last_report)
                self.export_csv_btn.config(state="normal" if has_report else "disabled")
                self.export_xlsx_btn.config(state="normal" if (has_report and OPENPYXL_AVAILABLE) else "disabled")
                self.copy_btn.config(state="normal")
                self.process_btn.set_enabled(True)
                self._add_history_entry(total_files - failures, total_kept, total_blocks, out_dir)
                self.show_toast((t('Готово: ') + f"{total_kept}" + t(' из ') + f"{total_blocks}" + t(' строк в ') + f"{total_files - failures}" + t(' файл(ах).')))

            self.after(0, finish)

        self.last_report = []
        threading.Thread(target=worker, daemon=True).start()

    def open_output_folder(self):
        if not self.last_out_dir:
            return
        try:
            os.startfile(self.last_out_dir)  # Windows
        except AttributeError:
            subprocess.Popen(["xdg-open", self.last_out_dir])

    # ---------- история ----------

    def _add_history_entry(self, files_count, kept, total, out_dir):
        entry = {
            "ts": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "files": files_count,
            "mode": t("Убрать") if self.mode_var.get() == MODE_REMOVE else t("Оставить"),
            "group": self.group_combo.get(),
            "kept": kept,
            "total": total,
            "out_dir": out_dir,
        }
        history = self.settings.setdefault("history", [])
        history.append(entry)
        self.settings["history"] = history[-50:]
        save_settings(self.settings)

    def show_history(self):
        win = tk.Toplevel(self)
        win.title(t("История обработок"))
        win.configure(bg=self.theme["panel"])
        win.geometry("560x360")
        win.transient(self)

        history = list(reversed(self.settings.get("history", [])))
        if not history:
            ttk.Label(win, text=t("Пока пусто — здесь появится история после первой обработки."),
                      style="Panel.TLabel", wraplength=500).pack(padx=16, pady=20)
            return

        box = tk.Listbox(win, bg=self.theme["panel2"], fg=self.theme["text"], relief="flat", font=FONT_MONO)
        box.pack(fill="both", expand=True, padx=12, pady=12)
        for h in history:
            box.insert("end", (f"{h['ts']}" + '  ·  ' + f"{h['files']}" + t(' файл(ов)  ·  ') + f"{h['mode']}" + ' (' + f"{h['group']}" + ')  ·  ' + f"{h['kept']}" + '/' + f"{h['total']}" + t(' строк')))

        def open_selected(event=None):
            sel = box.curselection()
            if not sel:
                return
            entry = history[sel[0]]
            if os.path.isdir(entry.get("out_dir", "")):
                os.startfile(entry["out_dir"])
            else:
                self.show_toast(t("Папка результата больше не существует."), kind="warn")

        box.bind("<Double-Button-1>", open_selected)
        ttk.Label(win, text=t("Двойной клик — открыть папку с результатом этого запуска."),
                  style="Panel.TLabel").pack(anchor="w", padx=12, pady=(0, 10))

    # ---------- экспорт ----------

    def export_csv(self):
        if not self.last_report:
            messagebox.showinfo(APP_TITLE, t("Сначала обработай файлы."))
            return
        path = filedialog.asksaveasfilename(
            title=t("Сохранить отчёт"), defaultextension=".csv",
            filetypes=[("CSV", "*.csv")], initialfile="report.csv",
        )
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow([t("Файл"), t("Слово"), t("Найдено раз")])
            for row in self.last_report:
                writer.writerow([row["file"], row["keyword"], row["count"]])
        self.show_toast((t('Отчёт сохранён: ') + f"{os.path.basename(path)}"))

    def export_xlsx(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showinfo(APP_TITLE, t("Для экспорта в Excel нужна библиотека openpyxl:\npip install openpyxl"))
            return
        if not self.last_report:
            messagebox.showinfo(APP_TITLE, t("Сначала обработай файлы."))
            return
        path = filedialog.asksaveasfilename(
            title=t("Сохранить отчёт"), defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")], initialfile="report.xlsx",
        )
        if not path:
            return
        self._write_xlsx_report(path, self.last_report, t("Отчёт"))
        self.show_toast((t('Отчёт сохранён: ') + f"{os.path.basename(path)}"))

    def export_all_groups(self):
        if not OPENPYXL_AVAILABLE:
            fmt = "csv"
        else:
            choice = self._ask_choice(t("Формат отчёта"), t("В каком формате сохранить:"), ["xlsx", "csv"], "xlsx")
            if not choice:
                return
            fmt = choice
        if not self.files:
            messagebox.showwarning(APP_TITLE, t("Сначала выбери файлы субтитров."))
            return

        rows = []
        for group_name, words in self.settings["groups"].items():
            if not words:
                continue
            try:
                matcher = self._build_matcher(words)
                singles = [(kw, build_single_matcher(kw, self.whole_word_var.get(),
                                                      self.case_sensitive_var.get(), self.regex_mode_var.get()))
                           for kw in words]
            except re.error:
                continue
            for path in self.files:
                try:
                    _, entries = peek_subtitle_blocks(path)
                except Exception:
                    continue
                all_text = "\n".join(e["text"] for e in entries)
                for kw, m in singles:
                    c = len(m.findall(all_text))
                    if c:
                        rows.append({"group": group_name, "file": os.path.basename(path), "keyword": kw, "count": c})

        if not rows:
            messagebox.showinfo(APP_TITLE, t("Совпадений не найдено ни в одной группе."))
            return

        default_name = f"report_all_groups.{fmt}"
        path = filedialog.asksaveasfilename(
            title=t("Сохранить отчёт по всем группам"), defaultextension=f".{fmt}",
            filetypes=[("Excel", "*.xlsx")] if fmt == "xlsx" else [("CSV", "*.csv")],
            initialfile=default_name,
        )
        if not path:
            return

        if fmt == "xlsx":
            self._write_xlsx_report(path, rows, t("Все группы"), with_group=True)
        else:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")
                writer.writerow([t("Группа"), t("Файл"), t("Слово"), t("Найдено раз")])
                for row in rows:
                    writer.writerow([row["group"], row["file"], row["keyword"], row["count"]])
        self.show_toast((t('Отчёт по всем группам сохранён: ') + f"{os.path.basename(path)}"))

    def _write_xlsx_report(self, path, rows, sheet_title, with_group=False):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]
        headers = ([t("Группа")] if with_group else []) + [t("Файл"), t("Слово"), t("Найдено раз")]
        ws.append(headers)
        for row in rows:
            values = ([row["group"]] if with_group else []) + [row["file"], row["keyword"], row["count"]]
            ws.append(values)

        totals = {}
        for row in rows:
            totals[row["keyword"]] = totals.get(row["keyword"], 0) + row["count"]
        if totals:
            chart_ws = wb.create_sheet(t("По словам"))
            chart_ws.append([t("Слово"), t("Всего")])
            for kw, count in sorted(totals.items(), key=lambda kv: kv[1], reverse=True):
                chart_ws.append([kw, count])
            chart = BarChart()
            chart.title = t("Найдено раз по словам")
            chart.y_axis.title = t("Раз")
            data = Reference(chart_ws, min_col=2, min_row=1, max_row=chart_ws.max_row)
            cats = Reference(chart_ws, min_col=1, min_row=2, max_row=chart_ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart_ws.add_chart(chart, "D2")

        wb.save(path)

    def copy_summary_to_clipboard(self):
        if not self.last_report:
            messagebox.showinfo(APP_TITLE, t("Сначала обработай файлы."))
            return
        totals = {}
        for row in self.last_report:
            totals[row["keyword"]] = totals.get(row["keyword"], 0) + row["count"]
        lines = [f"{kw}: {count}" for kw, count in sorted(totals.items(), key=lambda kv: kv[1], reverse=True)]
        text = "\n".join(lines)
        self.clipboard_clear()
        self.clipboard_append(text)
        self.show_toast(t("Сводка скопирована в буфер обмена."))

    # ---------- обновления ----------

    def check_updates(self):
        def worker():
            url = f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPO}/releases/latest"
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "subtitle-filter-app"})
                with urllib.request.urlopen(req, timeout=6) as resp:
                    data = json.loads(resp.read().decode("utf-8"))
                latest = data.get("tag_name", "").lstrip("v")
                html_url = data.get("html_url", "")
                asset_url = None
                for asset in data.get("assets", []):
                    if asset.get("name") == RELEASE_ASSET_NAME:
                        asset_url = asset.get("browser_download_url")
                        break
            except Exception as e:
                self.after(0, lambda: messagebox.showinfo(
                    APP_TITLE, (t('Не удалось проверить обновления.\n(') + f"{e}" + ')')
                ))
                return

            def show():
                if latest and latest != APP_VERSION:
                    self._offer_update(latest, html_url, asset_url)
                else:
                    messagebox.showinfo(APP_TITLE, t("У тебя установлена последняя версия."))

            self.after(0, show)

        threading.Thread(target=worker, daemon=True).start()

    def _offer_update(self, latest, html_url, asset_url):
        can_auto = bool(asset_url) and getattr(sys, "frozen", False)
        if can_auto:
            answer = messagebox.askyesnocancel(
                APP_TITLE,
                (t('Доступна новая версия: ') + f"{latest}" + t(' (у тебя ') + f"{APP_VERSION}"
                 + t(').\n\nДа — скачать и установить автоматически.\nНет — открыть страницу релиза в браузере.')),
            )
            if answer is None:
                return
            if answer:
                self._download_and_install(asset_url)
                return
            webbrowser.open(html_url)
        else:
            if messagebox.askyesno(
                APP_TITLE, (t('Доступна новая версия: ') + f"{latest}" + t(' (у тебя ') + f"{APP_VERSION}" + t(').\nОткрыть страницу релиза?'))
            ):
                webbrowser.open(html_url)

    def _download_and_install(self, asset_url):
        self.show_toast(t("Скачиваю обновление..."))

        def worker():
            try:
                tmp_dir = tempfile.mkdtemp(prefix="subtitlefilter_update_")
                new_exe = os.path.join(tmp_dir, RELEASE_ASSET_NAME)
                req = urllib.request.Request(asset_url, headers={"User-Agent": "subtitle-filter-app"})
                with urllib.request.urlopen(req, timeout=30) as resp, open(new_exe, "wb") as f:
                    shutil.copyfileobj(resp, f)

                current_exe = sys.executable
                bat_path = os.path.join(tmp_dir, "update.bat")
                with open(bat_path, "w", encoding="utf-8") as f:
                    f.write(
                        "@echo off\r\n"
                        ":wait\r\n"
                        "tasklist | findstr /i \"%s\" >nul\r\n"
                        "if not errorlevel 1 (\r\n"
                        "  timeout /t 1 /nobreak >nul\r\n"
                        "  goto wait\r\n"
                        ")\r\n"
                        "copy /y \"%s\" \"%s\" >nul\r\n"
                        "start \"\" \"%s\"\r\n"
                        "del \"%%~f0\"\r\n"
                        % (os.path.basename(current_exe), new_exe, current_exe, current_exe)
                    )

                def launch_and_exit():
                    subprocess.Popen(["cmd", "/c", bat_path], creationflags=0x00000008)
                    self._on_close()

                self.after(0, launch_and_exit)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror(APP_TITLE, (t('Не удалось установить обновление:\n') + f"{e}")))

        threading.Thread(target=worker, daemon=True).start()


def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    if "--cli" in sys.argv:
        sys.exit(run_cli(sys.argv[1:]))
    else:
        main()
